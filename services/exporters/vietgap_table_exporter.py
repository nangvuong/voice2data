"""
Module xuất dữ liệu ra file Excel (.xlsx), Word (.docx), PDF (.pdf)
theo đúng các biểu mẫu chuẩn của tài liệu VietGAP:
'vietGAP.pdf' (Phần II: Các biểu mẫu ghi chép lưu trữ hồ sơ):
- Biểu mẫu 1.1 & 1.2: Trang bìa & Các thông tin chung
- Biểu mẫu 1.5: Nhật ký mua hoặc sản xuất phân bón, thuốc BVTV và hóa chất
- Biểu mẫu 1.6: Nhật ký sử dụng phân bón và thuốc BVTV, hóa chất
- Biểu mẫu 1.7: Nhật ký thu gom, xử lý bao bì chứa đựng và thuốc BVTV dư thừa
- Biểu mẫu 1.8: Nhật ký thu hoạch và tiêu thụ sản phẩm
"""

import io
import os
import sys
import re
import logging
from datetime import datetime, date
from decimal import Decimal
from typing import Optional, Dict, Any, List, Tuple, Union

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib import pagesizes, colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    KeepTogether,
)

from services.database import get_db_connection
from services.exporters.cultivation_diary_exporter import _init_fonts, NumberedCanvas

logger = logging.getLogger("nong-ky-vietgap-table-exporter")


def _format_date(val: Any) -> str:
    """Định dạng ngày DD/MM/YYYY."""
    if not val:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%d/%m/%Y")
    s = str(val).strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"
    return s


def _format_qty(qty: Any, unit: Optional[str] = None) -> str:
    """Định dạng số lượng."""
    if qty is None or qty == "":
        return "-"
    try:
        val_f = float(qty)
        fmt = f"{int(val_f):,}" if val_f.is_integer() else f"{val_f:,.2f}"
        fmt = fmt.replace(",", "X").replace(".", ",").replace("X", ".")
        u = (unit or "").strip()
        return f"{fmt} {u}".strip() if u else fmt
    except (ValueError, TypeError):
        return str(qty)


def get_vietgap_full_tables_data(
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> Dict[str, Any]:
    """
    Truy vấn CSDL trích xuất toàn bộ dữ liệu theo các biểu mẫu chuẩn vietGAP.pdf:
    - cover: Trang bìa & thông tin chung (1.1 & 1.2)
    - b1_5: Bảng 1.5 Mua vật tư
    - b1_6: Bảng 1.6 Sử dụng phân & thuốc
    - b1_7: Bảng 1.7 Xử lý bao bì
    - b1_8: Bảng 1.8 Thu hoạch & Tiêu thụ
    """
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            # 1. Tìm season_id
            if not season_id and farmer_id:
                cur.execute(
                    """
                    SELECT cs.season_id 
                    FROM crop_seasons cs
                    JOIN plots p ON cs.plot_id = p.plot_id
                    JOIN farm f ON p.farm_id = f.farm_id
                    WHERE f.owner_user_id = %s
                    ORDER BY cs.season_id DESC LIMIT 1;
                    """,
                    (farmer_id,),
                )
                r = cur.fetchone()
                if r:
                    season_id = r["season_id"]

            if not season_id:
                cur.execute("SELECT season_id FROM crop_seasons ORDER BY season_id DESC LIMIT 1;")
                r = cur.fetchone()
                if r:
                    season_id = r["season_id"]

            if not season_id:
                raise ValueError("Không tìm thấy mùa vụ nào để xuất báo cáo VietGAP.")

            # 2. Thông tin chung (Biểu mẫu 1.1 & 1.2)
            cur.execute(
                """
                SELECT cs.season_id, cs.season_name, cs.crop_name, cs.variety AS season_variety,
                       cs.start_date, cs.expected_harvest_date, cs.expected_yield, cs.actual_yield, cs.yield_unit,
                       p.plot_id, p.name AS plot_name, p.plot_code, p.variety AS plot_variety,
                       p.area_size, p.area_unit, p.soil_type, p.water_source, p.land_use_history,
                       f.farm_id, f.name AS farm_name, f.farm_code, f.address AS farm_address,
                       f.certification_body, f.certification_no, f.certification_date, f.certification_expiry,
                       u.user_id, u.full_name AS farmer_name, u.phone AS farmer_phone
                FROM crop_seasons cs
                JOIN plots p ON cs.plot_id = p.plot_id
                JOIN farm f ON p.farm_id = f.farm_id
                LEFT JOIN users u ON f.owner_user_id = u.user_id
                WHERE cs.season_id = %s
                LIMIT 1;
                """,
                (season_id,),
            )
            cs_row = cur.fetchone() or {}

            variety_str = cs_row.get("season_variety") or cs_row.get("plot_variety") or cs_row.get("crop_name") or "-"
            if cs_row.get("crop_name") and cs_row.get("crop_name") not in variety_str:
                variety_str = f"{cs_row['crop_name']} ({variety_str})"

            cover_info = {
                "season_id": season_id,
                "farm_name": cs_row.get("farm_name") or "HỢP TÁC XÃ NÔNG NGHIỆP VIETGAP",
                "farmer_name": cs_row.get("farmer_name") or "Chủ hộ sản xuất",
                "address": cs_row.get("farm_address") or "Chưa cập nhật",
                "farm_code": cs_row.get("farm_code") or cs_row.get("plot_code") or "VN-GAP-001",
                "crop_season": cs_row.get("season_name") or f"Vụ #{season_id}",
                "variety": variety_str,
                "start_date": _format_date(cs_row.get("start_date")),
                "area_size": _format_qty(cs_row.get("area_size"), cs_row.get("area_unit") or "m²"),
                "soil_type": cs_row.get("soil_type") or "Đất đỏ bazan / phù sa",
                "water_source": cs_row.get("water_source") or "Giếng khoan tầng sâu / Nước sông suối tự nhiên",
                "land_use_history": cs_row.get("land_use_history") or "Đất nông nghiệp trồng trọt lâu năm",
                "certification_no": cs_row.get("certification_no") or "VietGAP-001",
                "certification_body": cs_row.get("certification_body") or "Chi cục Trồng trọt & BVTV",
            }

            # 3. Biểu mẫu 1.5: Nhật ký mua hoặc sản xuất phân bón, thuốc BVTV và hóa chất
            cur.execute(
                """
                SELECT ip.purchase_id, ip.purchase_date, ip.item_type,
                       COALESCE(p.name, f.name, m.name, 'Vật tư nông nghiệp') AS material_name,
                       ip.quantity, ip.unit,
                       s.name AS supplier_name, s.address AS supplier_address,
                       ip.expiry_date, ip.invoice_no, ip.note
                FROM input_purchases ip
                LEFT JOIN pesticides p ON ip.pesticide_id = p.pesticide_id
                LEFT JOIN fertilizers f ON ip.fertilizer_id = f.fertilizer_id
                LEFT JOIN materials m ON ip.material_id = m.material_id
                LEFT JOIN suppliers s ON ip.supplier_id = s.supplier_id
                WHERE ip.season_id = %s
                ORDER BY ip.purchase_date ASC, ip.purchase_id ASC;
                """,
                (season_id,),
            )
            raw_b1_5 = cur.fetchall()
            b1_5_list = []
            for idx, r in enumerate(raw_b1_5, start=1):
                b1_5_list.append(
                    {
                        "stt": idx,
                        "purchase_date": _format_date(r.get("purchase_date")),
                        "material_name": r.get("material_name") or "-",
                        "quantity": _format_qty(r.get("quantity"), r.get("unit")),
                        "supplier_name": r.get("supplier_name") or "-",
                        "supplier_address": r.get("supplier_address") or "-",
                        "expiry_date": _format_date(r.get("expiry_date")) if r.get("expiry_date") else "-",
                        "invoice_no": r.get("invoice_no") or r.get("note") or "-",
                    }
                )

            # 4. Biểu mẫu 1.6: Nhật ký sử dụng phân bón và thuốc BVTV, hóa chất
            cur.execute(
                """
                SELECT ca.activity_id, ca.activity_date, ca.activity_type, ca.growth_stage, ca.description,
                       f.name AS fert_name, fa.quantity AS fert_qty, fa.unit AS fert_unit, fa.application_method AS fert_method,
                       p.name AS pest_name, pa.dose_quantity, pa.dose_unit, pa.concentration_value, pa.concentration_unit,
                       pa.isolation_days, p.min_isolation_days, pa.target_pest
                FROM cultivation_activities ca
                LEFT JOIN fertilizer_applications fa ON ca.activity_id = fa.activity_id
                LEFT JOIN fertilizers f ON fa.fertilizer_id = f.fertilizer_id
                LEFT JOIN pesticide_applications pa ON ca.activity_id = pa.activity_id
                LEFT JOIN pesticides p ON pa.pesticide_id = p.pesticide_id
                WHERE ca.season_id = %s 
                  AND (ca.activity_type IN ('FERTILIZER', 'PESTICIDE') OR fa.application_id IS NOT NULL OR pa.application_id IS NOT NULL)
                ORDER BY ca.activity_date ASC, ca.activity_id ASC;
                """,
                (season_id,),
            )
            raw_b1_6 = cur.fetchall()
            b1_6_list = []
            for idx, r in enumerate(raw_b1_6, start=1):
                # Bón phân
                f_name = r.get("fert_name") or "-"
                f_qty_str = "-"
                if r.get("fert_qty") is not None:
                    f_qty_str = _format_qty(r["fert_qty"], r.get("fert_unit") or "kg")

                # Phun thuốc
                p_name = r.get("pest_name") or "-"
                p_dose_str = "-"
                if r.get("dose_quantity") is not None:
                    p_dose_str = _format_qty(r["dose_quantity"], r.get("dose_unit") or "lít")
                    if r.get("concentration_value") is not None:
                        p_dose_str += f" ({_format_qty(r['concentration_value'], r.get('concentration_unit') or '%')})"

                iso_str = "-"
                if r.get("isolation_days") is not None:
                    iso_str = f"{r['isolation_days']} ngày"
                elif r.get("min_isolation_days") is not None:
                    iso_str = f"{r['min_isolation_days']} ngày"
                elif p_name != "-":
                    iso_str = "7 ngày"

                b1_6_list.append(
                    {
                        "stt": idx,
                        "activity_date": _format_date(r.get("activity_date")),
                        "fert_name": f_name,
                        "fert_quantity": f_qty_str,
                        "pest_name": p_name,
                        "pest_dose": p_dose_str,
                        "isolation_days": iso_str,
                    }
                )

            # 5. Biểu mẫu 1.7: Nhật ký thu gom, xử lý bao bì và thuốc BVTV dư thừa
            cur.execute(
                """
                SELECT wd.disposal_id, wd.disposal_date, wd.waste_type, wd.quantity, wd.unit,
                       wd.storage_location, wd.treatment_method, wd.note
                FROM waste_disposals wd
                WHERE wd.season_id = %s
                ORDER BY wd.disposal_date ASC, wd.disposal_id ASC;
                """,
                (season_id,),
            )
            raw_b1_7 = cur.fetchall()
            b1_7_list = []
            for idx, r in enumerate(raw_b1_7, start=1):
                b1_7_list.append(
                    {
                        "stt": idx,
                        "disposal_date": _format_date(r.get("disposal_date")),
                        "waste_type": r.get("waste_type") or "Bao bì phân bón / chai lọ thuốc BVTV",
                        "storage_location": r.get("storage_location") or "Khu lưu chứa bao bì quy định",
                        "treatment_method": r.get("treatment_method") or "Thu gom xử lý theo quy định VietGAP",
                        "note": r.get("note") or "-",
                    }
                )

            # 6. Biểu mẫu 1.8: Nhật ký thu hoạch và tiêu thụ sản phẩm
            cur.execute(
                """
                SELECT hr.harvest_id, ca.activity_date AS harvest_date, hr.quantity AS harvest_qty, hr.unit AS harvest_unit,
                       hr.processing_location,
                       sr.sale_date, sr.buyer_name, sr.buyer_contact, sr.quantity AS sale_qty, sr.unit AS sale_unit
                FROM harvest_records hr
                LEFT JOIN cultivation_activities ca ON hr.activity_id = ca.activity_id
                LEFT JOIN sales_records sr ON hr.harvest_id = sr.harvest_id
                WHERE hr.season_id = %s
                ORDER BY ca.activity_date ASC, hr.harvest_id ASC;
                """,
                (season_id,),
            )
            raw_b1_8 = cur.fetchall()
            b1_8_list = []
            for idx, r in enumerate(raw_b1_8, start=1):
                buyer_info = r.get("buyer_name") or "-"
                if r.get("buyer_contact") and buyer_info != "-":
                    buyer_info += f" ({r['buyer_contact']})"

                b1_8_list.append(
                    {
                        "stt": idx,
                        "harvest_date": _format_date(r.get("harvest_date")),
                        "harvest_qty": _format_qty(r.get("harvest_qty"), r.get("harvest_unit") or "kg"),
                        "processing_location": r.get("processing_location") or "Sơ chế tại vườn",
                        "sale_date": _format_date(r.get("sale_date")) if r.get("sale_date") else "-",
                        "buyer_info": buyer_info,
                        "sale_qty": _format_qty(r.get("sale_qty"), r.get("sale_unit") or "kg") if r.get("sale_qty") is not None else "-",
                    }
                )

            return {
                "season_id": season_id,
                "cover_info": cover_info,
                "b1_5_purchases": b1_5_list,
                "b1_6_usages": b1_6_list,
                "b1_7_waste": b1_7_list,
                "b1_8_harvest_sales": b1_8_list,
                "exported_at": datetime.now().strftime("%d/%m/%Y %H:%M"),
            }
    finally:
        conn.close()


# ==============================================================================
# 1. XUẤT EXCEL THEO CÁC BIỂU MẪU VIETGAP
# ==============================================================================

def build_vietgap_tables_workbook(data: Dict[str, Any]) -> openpyxl.Workbook:
    """
    Tạo workbook Excel chứa đầy đủ các sheet biểu mẫu VietGAP chuẩn:
    - Sheet 1: 1.1 & 1.2 Thông tin chung
    - Sheet 2: Biểu mẫu 1.5 - Mua vật tư
    - Sheet 3: Biểu mẫu 1.6 - Sử dụng phân & thuốc (Header 2 tầng)
    - Sheet 4: Biểu mẫu 1.7 - Xử lý bao bì
    - Sheet 5: Biểu mẫu 1.8 - Thu hoạch & Tiêu thụ
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Xóa sheet mặc định

    font_title = Font(name="Arial", size=14, bold=True, color="1B5E20")
    font_sub = Font(name="Arial", size=10, italic=True, color="555555")
    font_sec = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_th = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_lbl = Font(name="Arial", size=10, bold=True, color="1B5E20")
    font_val = Font(name="Arial", size=10, color="111111")
    font_td = Font(name="Arial", size=9.5, color="212121")

    fill_sec = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
    fill_th = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
    fill_lbl = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    fill_zebra = PatternFill(start_color="F9FBF7", end_color="F9FBF7", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    cover = data.get("cover_info", {})

    # ---------- SHEET 1: THÔNG TIN CHUNG (1.1 & 1.2) ----------
    ws1 = wb.create_sheet(title="Thông Tin Chung (1.1-1.2)")
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:F1")
    ws1["A1"].value = cover.get("farm_name", "").upper()
    ws1["A1"].font = Font(name="Arial", size=10, bold=True, color="388E3C")
    ws1["A1"].alignment = align_center

    ws1.merge_cells("A2:F2")
    ws1["A2"].value = "BIỂU MẪU 1: NHẬT KÝ SẢN XUẤT THEO TIÊU CHUẨN VIETGAP"
    ws1["A2"].font = font_title
    ws1["A2"].alignment = align_center

    ws1.merge_cells("A4:F4")
    ws1["A4"].value = "1.2. CÁC THÔNG TIN CHUNG VỀ NÔNG HỘ VÀ VƯỜN TRỒNG"
    ws1["A4"].font = font_sec
    ws1["A4"].fill = fill_sec
    ws1["A4"].alignment = align_left

    general_rows = [
        ("Tên hộ sản xuất / Nông hộ:", cover.get("farmer_name", "-")),
        ("Địa chỉ nông hộ:", cover.get("address", "-")),
        ("Mã số vùng trồng / P.U.C:", cover.get("farm_code", "-")),
        ("Vụ canh tác:", cover.get("crop_season", "-")),
        ("Tên giống cây trồng:", cover.get("variety", "-")),
        ("Thời gian bắt đầu trồng:", cover.get("start_date", "-")),
        ("Diện tích vườn áp dụng VietGAP:", cover.get("area_size", "-")),
        ("Loại đất canh tác:", cover.get("soil_type", "-")),
        ("Nguồn nước tưới:", cover.get("water_source", "-")),
        ("Lịch sử khu đất canh tác:", cover.get("land_use_history", "-")),
        ("Số chứng nhận VietGAP:", cover.get("certification_no", "-")),
        ("Đơn vị chứng nhận:", cover.get("certification_body", "-")),
    ]

    r_idx = 5
    for lbl, val in general_rows:
        ws1.row_dimensions[r_idx].height = 22
        ws1.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=2)
        c_lbl = ws1.cell(r_idx, 1, value=lbl)
        c_lbl.font = font_lbl
        c_lbl.fill = fill_lbl
        c_lbl.alignment = align_left

        ws1.merge_cells(start_row=r_idx, start_column=3, end_row=r_idx, end_column=6)
        c_val = ws1.cell(r_idx, 3, value=val)
        c_val.font = font_val
        c_val.fill = fill_white
        c_val.alignment = align_left

        for c in range(1, 7):
            ws1.cell(r_idx, c).border = thin_border
        r_idx += 1

    ws1.column_dimensions["A"].width = 15
    ws1.column_dimensions["B"].width = 25
    ws1.column_dimensions["C"].width = 25
    ws1.column_dimensions["D"].width = 20
    ws1.column_dimensions["E"].width = 20
    ws1.column_dimensions["F"].width = 20

    # ---------- SHEET 2: BẢNG 1.5 MUA VẬT TƯ ----------
    ws2 = wb.create_sheet(title="Biểu Mẫu 1.5 (Mua Vật Tư)")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:G1")
    ws2["A1"].value = "1.5. NHẬT KÝ MUA HOẶC SẢN XUẤT PHÂN BÓN, THUỐC BVTV VÀ HÓA CHẤT"
    ws2["A1"].font = font_title
    ws2["A1"].alignment = align_center

    th_1_5 = [
        ("STT", 6, align_center),
        ("Ngày mua hoặc sản xuất", 18, align_center),
        ("Tên vật tư", 30, align_left),
        ("Khối lượng mua", 18, align_right),
        ("Nơi cung cấp (Tên đại lý)", 26, align_left),
        ("Địa chỉ nơi cung cấp", 26, align_left),
        ("Hạn sử dụng", 16, align_center),
    ]

    ws2.row_dimensions[3].height = 28
    for c_idx, (h_name, w, a) in enumerate(th_1_5, start=1):
        c = ws2.cell(3, c_idx, value=h_name)
        c.font = font_th
        c.fill = fill_th
        c.alignment = a
        c.border = thin_border
        ws2.column_dimensions[get_column_letter(c_idx)].width = w

    r_idx = 4
    for item in data.get("b1_5_purchases", []):
        fill_r = fill_zebra if item["stt"] % 2 == 0 else fill_white
        ws2.row_dimensions[r_idx].height = 22
        ws2.cell(r_idx, 1, item["stt"]).alignment = align_center
        ws2.cell(r_idx, 2, item["purchase_date"]).alignment = align_center
        ws2.cell(r_idx, 3, item["material_name"]).alignment = align_left
        ws2.cell(r_idx, 4, item["quantity"]).alignment = align_right
        ws2.cell(r_idx, 5, item["supplier_name"]).alignment = align_left
        ws2.cell(r_idx, 6, item["supplier_address"]).alignment = align_left
        ws2.cell(r_idx, 7, item["expiry_date"]).alignment = align_center

        for c in range(1, 8):
            cell = ws2.cell(r_idx, c)
            cell.font = font_td
            cell.fill = fill_r
            cell.border = thin_border
        r_idx += 1

    # ---------- SHEET 3: BẢNG 1.6 SỬ DỤNG PHÂN & THUỐC (HEADER 2 TẦNG) ----------
    ws3 = wb.create_sheet(title="Biểu Mẫu 1.6 (Sử Dụng)")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:G1")
    ws3["A1"].value = "1.6. NHẬT KÝ SỬ DỤNG PHÂN BÓN VÀ THUỐC BVTV, HÓA CHẤT"
    ws3["A1"].font = font_title
    ws3["A1"].alignment = align_center

    # Header tầng 1 (Dòng 3)
    ws3.row_dimensions[3].height = 24
    ws3.row_dimensions[4].height = 24

    ws3.merge_cells("A3:A4")
    ws3["A3"].value = "STT"
    ws3.merge_cells("B3:B4")
    ws3["B3"].value = "Ngày, tháng, năm"

    ws3.merge_cells("C3:D3")
    ws3["C3"].value = "Bón phân"

    ws3.merge_cells("E3:G3")
    ws3["E3"].value = "Sử dụng thuốc BVTV"

    # Header tầng 2 (Dòng 4)
    ws3["C4"].value = "Tên phân bón"
    ws3["D4"].value = "Lượng sử dụng"
    ws3["E4"].value = "Tên thuốc"
    ws3["F4"].value = "Nồng độ và lượng sử dụng"
    ws3["G4"].value = "Thời gian cách ly"

    for r_h in (3, 4):
        for col in range(1, 8):
            c = ws3.cell(r_h, col)
            c.font = font_th
            c.fill = fill_th
            c.alignment = align_center
            c.border = thin_border

    widths_1_6 = [6, 16, 26, 18, 24, 26, 18]
    for idx, w in enumerate(widths_1_6, start=1):
        ws3.column_dimensions[get_column_letter(idx)].width = w

    r_idx = 5
    for item in data.get("b1_6_usages", []):
        fill_r = fill_zebra if item["stt"] % 2 == 0 else fill_white
        ws3.row_dimensions[r_idx].height = 22
        ws3.cell(r_idx, 1, item["stt"]).alignment = align_center
        ws3.cell(r_idx, 2, item["activity_date"]).alignment = align_center
        ws3.cell(r_idx, 3, item["fert_name"]).alignment = align_left
        ws3.cell(r_idx, 4, item["fert_quantity"]).alignment = align_right
        ws3.cell(r_idx, 5, item["pest_name"]).alignment = align_left
        ws3.cell(r_idx, 6, item["pest_dose"]).alignment = align_right
        ws3.cell(r_idx, 7, item["isolation_days"]).alignment = align_center

        for c in range(1, 8):
            cell = ws3.cell(r_idx, c)
            cell.font = font_td
            cell.fill = fill_r
            cell.border = thin_border
        r_idx += 1

    # ---------- SHEET 4: BẢNG 1.7 XỬ LÝ BAO BÌ ----------
    ws4 = wb.create_sheet(title="Biểu Mẫu 1.7 (Bao Bì)")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:E1")
    ws4["A1"].value = "1.7. NHẬT KÝ THU GOM, XỬ LÝ BAO BÌ CHỨA ĐỰNG VÀ THUỐC BVTV DƯ THỪA"
    ws4["A1"].font = font_title
    ws4["A1"].alignment = align_center

    th_1_7 = [
        ("STT", 6, align_center),
        ("Ngày, tháng, năm", 18, align_center),
        ("Loại bao bì, thùng chứa, thuốc dư thừa", 32, align_left),
        ("Nơi tồn trữ, hủy bỏ", 28, align_left),
        ("Cách xử lý", 32, align_left),
    ]

    ws4.row_dimensions[3].height = 28
    for c_idx, (h_name, w, a) in enumerate(th_1_7, start=1):
        c = ws4.cell(3, c_idx, value=h_name)
        c.font = font_th
        c.fill = fill_th
        c.alignment = a
        c.border = thin_border
        ws4.column_dimensions[get_column_letter(c_idx)].width = w

    r_idx = 4
    for item in data.get("b1_7_waste", []):
        fill_r = fill_zebra if item["stt"] % 2 == 0 else fill_white
        ws4.row_dimensions[r_idx].height = 22
        ws4.cell(r_idx, 1, item["stt"]).alignment = align_center
        ws4.cell(r_idx, 2, item["disposal_date"]).alignment = align_center
        ws4.cell(r_idx, 3, item["waste_type"]).alignment = align_left
        ws4.cell(r_idx, 4, item["storage_location"]).alignment = align_left
        ws4.cell(r_idx, 5, item["treatment_method"]).alignment = align_left

        for c in range(1, 6):
            cell = ws4.cell(r_idx, c)
            cell.font = font_td
            cell.fill = fill_r
            cell.border = thin_border
        r_idx += 1

    # ---------- SHEET 5: BẢNG 1.8 THU HOẠCH & TIÊU THỤ ----------
    ws5 = wb.create_sheet(title="Biểu Mẫu 1.8 (Thu Hoạch)")
    ws5.views.sheetView[0].showGridLines = True

    ws5.merge_cells("A1:G1")
    ws5["A1"].value = "1.8. NHẬT KÝ THU HOẠCH VÀ TIÊU THỤ SẢN PHẨM"
    ws5["A1"].font = font_title
    ws5["A1"].alignment = align_center

    th_1_8 = [
        ("STT", 6, align_center),
        ("Ngày thu hoạch", 16, align_center),
        ("Sản lượng thu hoạch", 20, align_right),
        ("Địa điểm, cách sơ chế", 28, align_left),
        ("Thời gian xuất bán", 18, align_center),
        ("Cơ sở thu mua / tiêu thụ", 28, align_left),
        ("Khối lượng tiêu thụ", 20, align_right),
    ]

    ws5.row_dimensions[3].height = 28
    for c_idx, (h_name, w, a) in enumerate(th_1_8, start=1):
        c = ws5.cell(3, c_idx, value=h_name)
        c.font = font_th
        c.fill = fill_th
        c.alignment = a
        c.border = thin_border
        ws5.column_dimensions[get_column_letter(c_idx)].width = w

    r_idx = 4
    for item in data.get("b1_8_harvest_sales", []):
        fill_r = fill_zebra if item["stt"] % 2 == 0 else fill_white
        ws5.row_dimensions[r_idx].height = 22
        ws5.cell(r_idx, 1, item["stt"]).alignment = align_center
        ws5.cell(r_idx, 2, item["harvest_date"]).alignment = align_center
        ws5.cell(r_idx, 3, item["harvest_qty"]).alignment = align_right
        ws5.cell(r_idx, 4, item["processing_location"]).alignment = align_left
        ws5.cell(r_idx, 5, item["sale_date"]).alignment = align_center
        ws5.cell(r_idx, 6, item["buyer_info"]).alignment = align_left
        ws5.cell(r_idx, 7, item["sale_qty"]).alignment = align_right

        for c in range(1, 8):
            cell = ws5.cell(r_idx, c)
            cell.font = font_td
            cell.fill = fill_r
            cell.border = thin_border
        r_idx += 1

    return wb


def export_vietgap_tables_excel_bytes(
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> Tuple[io.BytesIO, str]:
    """Xuất các biểu mẫu VietGAP ra file Excel theo dạng buffer."""
    data = get_vietgap_full_tables_data(season_id=season_id, farmer_id=farmer_id)
    wb = build_vietgap_tables_workbook(data)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    season_name = data.get("cover_info", {}).get("crop_season", "VietGAP")
    safe = re.sub(r"[^\w\-_]+", "_", season_name).strip("_")
    date_str = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"So_Tay_VietGAP_Cac_Bang_{safe}_{date_str}.xlsx"

    return buf, filename


def export_vietgap_tables_excel_file(
    output_path: str,
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> str:
    """Xuất các biểu mẫu VietGAP ra file Excel trên đĩa cứng."""
    data = get_vietgap_full_tables_data(season_id=season_id, farmer_id=farmer_id)
    wb = build_vietgap_tables_workbook(data)

    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    wb.save(output_path)
    logger.info("Đã xuất sổ tay VietGAP ra Excel thành công: %s", output_path)
    return output_path


# ==============================================================================
# 2. XUẤT WORD DOCX THEO CÁC BIỂU MẪU VIETGAP
# ==============================================================================

def _apply_docx_table_style(table, col_widths):
    """Áp dụng viền và độ rộng cho bảng Word."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    border_spec = {"sz": 4, "val": "single", "color": "999999"}
    cell_borders = {
        "top": border_spec,
        "bottom": border_spec,
        "left": border_spec,
        "right": border_spec,
    }

    for row_idx, row in enumerate(table.rows):
        is_header = (row_idx == 0)
        # Chống tách dòng
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))

        for c_idx, (cell, w) in enumerate(zip(row.cells, col_widths)):
            cell.width = w
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = tcPr.first_child_found_in("w:tcBorders")
            if tcBorders is None:
                tcBorders = OxmlElement('w:tcBorders')
                tcPr.append(tcBorders)
            for edge, edge_data in cell_borders.items():
                element = OxmlElement(f'w:{edge}')
                for key, val in edge_data.items():
                    element.set(qn(f'w:{key}'), str(val))
                tcBorders.append(element)

            if is_header:
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8F5E9"/>')
                tcPr.append(shd)
            elif row_idx % 2 == 0:
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F9FBF7"/>')
                tcPr.append(shd)


def build_vietgap_tables_docx(data: Dict[str, Any]) -> docx.Document:
    """Tạo Document Word (.docx) chứa tất cả các biểu mẫu VietGAP."""
    doc = docx.Document()

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.27)
    section.right_margin = Cm(1.27)

    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(11)

    cover = data.get("cover_info", {})

    # ---------- PHẦN II: TIÊU ĐỀ HỒ SƠ VIETGAP ----------
    p_main = doc.add_paragraph()
    p_main.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_m = p_main.add_run("PHẦN II: CÁC BIỂU MẪU GHI CHÉP LƯU TRỮ HỒ SƠ VIETGAP\nSỔ TAY HƯỚNG DẪN KỸ THUẬT CANH TÁC THEO TIÊU CHUẨN VIETGAP")
    r_m.font.name = "Times New Roman"
    r_m.font.size = Pt(14)
    r_m.font.bold = True
    r_m.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)

    # 1.1 & 1.2: THÔNG TIN CHUNG
    p_sec1 = doc.add_paragraph()
    p_sec1.paragraph_format.space_before = Pt(8)
    p_sec1.paragraph_format.space_after = Pt(4)
    r_s1 = p_sec1.add_run("1. BIỂU MẪU 1: NHẬT KÝ SẢN XUẤT (1.1 TRANG BÌA & 1.2 CÁC THÔNG TIN CHUNG)")
    r_s1.font.bold = True
    r_s1.font.size = Pt(12)
    r_s1.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)

    gen_items = [
        ("Tổ hợp tác / Hợp tác xã:", cover.get("farm_name", "-")),
        ("Tên nông hộ / Chủ vườn:", cover.get("farmer_name", "-")),
        ("Địa chỉ:", cover.get("address", "-")),
        ("Mã số vùng trồng (P.U.C):", cover.get("farm_code", "-")),
        ("Tên giống cây trồng:", cover.get("variety", "-")),
        ("Thời gian trồng:", cover.get("start_date", "-")),
        ("Diện tích vườn áp dụng VietGAP:", cover.get("area_size", "-")),
        ("Nguồn nước tưới:", cover.get("water_source", "-")),
        ("Lịch sử khu đất canh tác:", cover.get("land_use_history", "-")),
    ]

    for lbl, val in gen_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"- {lbl} ")
        r1.font.bold = True
        r2 = p.add_run(str(val))
        r2.font.bold = False

    # ---------- BIỂU MẪU 1.5: MUA VẬT TƯ ----------
    p_15 = doc.add_paragraph()
    p_15.paragraph_format.space_before = Pt(14)
    p_15.paragraph_format.space_after = Pt(4)
    r15 = p_15.add_run("1.5. NHẬT KÝ: MUA HOẶC SẢN XUẤT PHÂN BÓN, THUỐC BVTV VÀ HÓA CHẤT")
    r15.font.bold = True
    r15.font.size = Pt(11)
    r15.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)

    rows_15 = data.get("b1_5_purchases", [])
    t15 = doc.add_table(rows=len(rows_15) + 1, cols=7)
    w15 = [Cm(1.2), Cm(3.2), Cm(6.5), Cm(3.2), Cm(5.0), Cm(5.0), Cm(3.0)]

    headers_15 = [
        "STT",
        "Ngày mua/sản xuất",
        "Tên vật tư",
        "Khối lượng mua",
        "Nơi cung cấp (Đại lý)",
        "Địa chỉ nơi cung cấp",
        "Hạn sử dụng",
    ]
    for idx, h in enumerate(headers_15):
        cell = t15.cell(0, idx)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)

    for r_idx, item in enumerate(rows_15, start=1):
        vals = [
            str(item["stt"]),
            item["purchase_date"],
            item["material_name"],
            item["quantity"],
            item["supplier_name"],
            item["supplier_address"],
            item["expiry_date"],
        ]
        for c_idx, v in enumerate(vals):
            cell = t15.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in (0, 1, 6) else (WD_ALIGN_PARAGRAPH.RIGHT if c_idx == 3 else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(v)
            r.font.size = Pt(9)

    _apply_docx_table_style(t15, w15)

    # ---------- BIỂU MẪU 1.6: SỬ DỤNG PHÂN & THUỐC ----------
    p_16 = doc.add_paragraph()
    p_16.paragraph_format.space_before = Pt(14)
    p_16.paragraph_format.space_after = Pt(4)
    r16 = p_16.add_run("1.6. NHẬT KÝ: SỬ DỤNG PHÂN BÓN VÀ THUỐC BVTV, HÓA CHẤT")
    r16.font.bold = True
    r16.font.size = Pt(11)
    r16.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)

    rows_16 = data.get("b1_6_usages", [])
    # 2 dòng tiêu đề + n dòng dữ liệu
    t16 = doc.add_table(rows=len(rows_16) + 2, cols=7)
    w16 = [Cm(1.2), Cm(3.2), Cm(5.5), Cm(3.6), Cm(5.5), Cm(4.8), Cm(3.3)]

    # Dòng tiêu đề 1: merge Bón phân (cols 2-3) và Phun thuốc (cols 4-6)
    cell_stt = t16.cell(0, 0)
    cell_stt.merge(t16.cell(1, 0))
    cell_stt.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cell_stt.paragraphs[0].add_run("STT")
    r.font.bold = True
    r.font.size = Pt(9.5)

    cell_date = t16.cell(0, 1)
    cell_date.merge(t16.cell(1, 1))
    cell_date.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cell_date.paragraphs[0].add_run("Ngày, tháng, năm")
    r.font.bold = True
    r.font.size = Pt(9.5)

    cell_fert = t16.cell(0, 2)
    cell_fert.merge(t16.cell(0, 3))
    cell_fert.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cell_fert.paragraphs[0].add_run("Bón phân")
    r.font.bold = True
    r.font.size = Pt(9.5)

    cell_pest = t16.cell(0, 4)
    cell_pest.merge(t16.cell(0, 6))
    cell_pest.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cell_pest.paragraphs[0].add_run("Sử dụng thuốc BVTV")
    r.font.bold = True
    r.font.size = Pt(9.5)

    # Dòng tiêu đề 2
    sub_headers_16 = [
        (2, "Tên phân bón"),
        (3, "Lượng sử dụng"),
        (4, "Tên thuốc"),
        (5, "Nồng độ & Lượng dùng"),
        (6, "Thời gian cách ly"),
    ]
    for c_idx, sub_t in sub_headers_16:
        c = t16.cell(1, c_idx)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.paragraphs[0].add_run(sub_t)
        r.font.bold = True
        r.font.size = Pt(9)

    for r_idx, item in enumerate(rows_16, start=2):
        vals = [
            str(item["stt"]),
            item["activity_date"],
            item["fert_name"],
            item["fert_quantity"],
            item["pest_name"],
            item["pest_dose"],
            item["isolation_days"],
        ]
        for c_idx, v in enumerate(vals):
            cell = t16.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in (0, 1, 6) else (WD_ALIGN_PARAGRAPH.RIGHT if c_idx in (3, 5) else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(v)
            r.font.size = Pt(9)

    _apply_docx_table_style(t16, w16)

    # ---------- BIỂU MẪU 1.7: XỬ LÝ BAO BÌ ----------
    p_17 = doc.add_paragraph()
    p_17.paragraph_format.space_before = Pt(14)
    p_17.paragraph_format.space_after = Pt(4)
    r17 = p_17.add_run("1.7. NHẬT KÝ: THU GOM, XỬ LÝ BAO BÌ CHỨA ĐỰNG VÀ THUỐC BVTV DƯ THỪA")
    r17.font.bold = True
    r17.font.size = Pt(11)
    r17.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)

    rows_17 = data.get("b1_7_waste", [])
    t17 = doc.add_table(rows=len(rows_17) + 1, cols=5)
    w17 = [Cm(1.2), Cm(3.5), Cm(8.0), Cm(6.5), Cm(7.9)]

    headers_17 = [
        "STT",
        "Ngày, tháng, năm",
        "Loại bao bì, thùng chứa, thuốc dư thừa",
        "Nơi tồn trữ, hủy bỏ",
        "Cách xử lý",
    ]
    for idx, h in enumerate(headers_17):
        cell = t17.cell(0, idx)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)

    for r_idx, item in enumerate(rows_17, start=1):
        vals = [
            str(item["stt"]),
            item["disposal_date"],
            item["waste_type"],
            item["storage_location"],
            item["treatment_method"],
        ]
        for c_idx, v in enumerate(vals):
            cell = t17.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in (0, 1) else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(v)
            r.font.size = Pt(9)

    _apply_docx_table_style(t17, w17)

    # ---------- BIỂU MẪU 1.8: THU HOẠCH & TIÊU THỤ ----------
    p_18 = doc.add_paragraph()
    p_18.paragraph_format.space_before = Pt(14)
    p_18.paragraph_format.space_after = Pt(4)
    r18 = p_18.add_run("1.8. NHẬT KÝ: THU HOẠCH VÀ TIÊU THỤ SẢN PHẨM")
    r18.font.bold = True
    r18.font.size = Pt(11)
    r18.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)

    rows_18 = data.get("b1_8_harvest_sales", [])
    t18 = doc.add_table(rows=len(rows_18) + 1, cols=7)
    w18 = [Cm(1.2), Cm(3.2), Cm(3.8), Cm(6.2), Cm(3.2), Cm(6.0), Cm(3.5)]

    headers_18 = [
        "STT",
        "Ngày thu hoạch",
        "Sản lượng (kg)",
        "Địa điểm, cách sơ chế",
        "Ngày xuất bán",
        "Cơ sở thu mua / tiêu thụ",
        "Khối lượng bán",
    ]
    for idx, h in enumerate(headers_18):
        cell = t18.cell(0, idx)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)

    for r_idx, item in enumerate(rows_18, start=1):
        vals = [
            str(item["stt"]),
            item["harvest_date"],
            item["harvest_qty"],
            item["processing_location"],
            item["sale_date"],
            item["buyer_info"],
            item["sale_qty"],
        ]
        for c_idx, v in enumerate(vals):
            cell = t18.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in (0, 1, 4) else (WD_ALIGN_PARAGRAPH.RIGHT if c_idx in (2, 6) else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(v)
            r.font.size = Pt(9)

    _apply_docx_table_style(t18, w18)

    # ---------- CHỮ KÝ XÁC NHẬN ----------
    now = datetime.now()
    p_sig_date = doc.add_paragraph()
    p_sig_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig_date.paragraph_format.space_before = Pt(16)
    p_sig_date.paragraph_format.space_after = Pt(4)
    r_sd = p_sig_date.add_run(f"Ngày {now.day:02d} tháng {now.month:02d} năm {now.year}")
    r_sd.font.italic = True

    sig_t = doc.add_table(rows=3, cols=2)
    sig_t.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_t.autofit = False
    for r in sig_t.rows:
        r.cells[0].width = Cm(13.5)
        r.cells[1].width = Cm(13.5)

    sig_t.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sig_t.cell(0, 0).paragraphs[0].add_run("CHỦ NÔNG HỘ / NGƯỜI GHI CHÉP")
    r.font.bold = True
    sig_t.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sig_t.cell(0, 1).paragraphs[0].add_run("CÁN BỘ GIÁM SÁT VIETGAP")
    r.font.bold = True

    sig_t.cell(1, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sig_t.cell(1, 0).paragraphs[0].add_run("(Ký và ghi rõ họ tên)")
    r.font.italic = True
    r.font.size = Pt(9.5)
    sig_t.cell(1, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sig_t.cell(1, 1).paragraphs[0].add_run("(Ký, đóng dấu xác nhận)")
    r.font.italic = True
    r.font.size = Pt(9.5)

    sig_t.cell(2, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig_t.cell(2, 0).paragraphs[0].paragraph_format.space_before = Pt(40)
    r = sig_t.cell(2, 0).paragraphs[0].add_run(cover.get("farmer_name", ""))
    r.font.bold = True

    return doc


def export_vietgap_tables_docx_bytes(
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> Tuple[io.BytesIO, str]:
    """Xuất các biểu mẫu VietGAP ra file Word (.docx) dạng buffer."""
    data = get_vietgap_full_tables_data(season_id=season_id, farmer_id=farmer_id)
    doc = build_vietgap_tables_docx(data)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    season_name = data.get("cover_info", {}).get("crop_season", "VietGAP")
    safe = re.sub(r"[^\w\-_]+", "_", season_name).strip("_")
    date_str = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"So_Tay_VietGAP_Cac_Bang_{safe}_{date_str}.docx"

    return buf, filename


def export_vietgap_tables_docx_file(
    output_path: str,
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> str:
    """Xuất các biểu mẫu VietGAP ra file Word (.docx) trên đĩa cứng."""
    data = get_vietgap_full_tables_data(season_id=season_id, farmer_id=farmer_id)
    doc = build_vietgap_tables_docx(data)

    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    doc.save(output_path)
    logger.info("Đã xuất sổ tay VietGAP ra DOCX thành công: %s", output_path)
    return output_path


# ==============================================================================
# 3. XUẤT PDF THEO CÁC BIỂU MẪU VIETGAP
# ==============================================================================

def build_vietgap_tables_pdf(data: Dict[str, Any], output_target: Union[str, io.BytesIO]):
    """Tạo file PDF chứa toàn bộ các biểu mẫu VietGAP chuẩn."""
    _init_fonts()

    pagesize = pagesizes.landscape(pagesizes.A4)
    margin = 36
    printable_width = pagesize[0] - 2 * margin  # 769.89 pt

    doc = SimpleDocTemplate(
        output_target,
        pagesize=pagesize,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
    )

    styles = getSampleStyleSheet()

    style_title = ParagraphStyle(
        "VGTitle",
        fontName="VN-Times-Bold",
        fontSize=14,
        leading=17,
        alignment=1,
        textColor=colors.HexColor("#1B5E20"),
        spaceAfter=8,
    )

    style_sec = ParagraphStyle(
        "VGSec",
        fontName="VN-Times-Bold",
        fontSize=10.5,
        leading=13,
        textColor=colors.HexColor("#1B5E20"),
        spaceBefore=10,
        spaceAfter=4,
    )

    style_meta = ParagraphStyle(
        "VGMeta",
        fontName="VN-Times",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#222222"),
    )

    style_th = ParagraphStyle(
        "VGTH",
        fontName="VN-Times-Bold",
        fontSize=8.5,
        leading=10,
        alignment=1,
        textColor=colors.HexColor("#1B5E20"),
    )

    style_td = ParagraphStyle(
        "VGTD",
        fontName="VN-Times",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#212121"),
    )

    style_td_center = ParagraphStyle(
        "VGTDCenter",
        parent=style_td,
        alignment=1,
    )

    style_td_right = ParagraphStyle(
        "VGTDRight",
        parent=style_td,
        alignment=2,
    )

    cover = data.get("cover_info", {})
    elements = []

    # 1. TIÊU ĐỀ CHÍNH
    elements.append(
        Paragraph("PHẦN II: CÁC BIỂU MẪU GHI CHÉP LƯU TRỮ HỒ SƠ VIETGAP", style_title)
    )

    # 2. THÔNG TIN CHUNG (Biểu mẫu 1.1 & 1.2)
    elements.append(Paragraph("1. BIỂU MẪU 1: NHẬT KÝ SẢN XUẤT (1.1 TRANG BÌA & 1.2 CÁC THÔNG TIN CHUNG)", style_sec))

    meta_grid = [
        [
            Paragraph(f"<b>Tổ hợp tác / HTX:</b> {cover.get('farm_name', '-')}", style_meta),
            Paragraph(f"<b>Vụ canh tác:</b> {cover.get('crop_season', '-')}", style_meta),
        ],
        [
            Paragraph(f"<b>Tên nông hộ:</b> {cover.get('farmer_name', '-')}", style_meta),
            Paragraph(f"<b>Giống cây trồng:</b> {cover.get('variety', '-')}", style_meta),
        ],
        [
            Paragraph(f"<b>Địa chỉ:</b> {cover.get('address', '-')}", style_meta),
            Paragraph(f"<b>Thời gian bắt đầu trồng:</b> {cover.get('start_date', '-')}", style_meta),
        ],
        [
            Paragraph(f"<b>Mã số vùng trồng (P.U.C):</b> {cover.get('farm_code', '-')}", style_meta),
            Paragraph(f"<b>Diện tích VietGAP:</b> {cover.get('area_size', '-')}", style_meta),
        ],
        [
            Paragraph(f"<b>Nguồn nước tưới:</b> {cover.get('water_source', '-')}", style_meta),
            Paragraph(f"<b>Loại đất:</b> {cover.get('soil_type', '-')}", style_meta),
        ],
    ]
    t_meta = Table(meta_grid, colWidths=[printable_width * 0.52, printable_width * 0.48])
    t_meta.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 1),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
            ]
        )
    )
    elements.append(t_meta)

    # 3. BIỂU MẪU 1.5: MUA VẬT TƯ
    elements.append(Paragraph("1.5. NHẬT KÝ: MUA HOẶC SẢN XUẤT PHÂN BÓN, THUỐC BVTV VÀ HÓA CHẤT", style_sec))
    t15_data = [
        [
            Paragraph("STT", style_th),
            Paragraph("Ngày mua / sản xuất", style_th),
            Paragraph("Tên vật tư", style_th),
            Paragraph("Khối lượng mua", style_th),
            Paragraph("Nơi cung cấp (Đại lý)", style_th),
            Paragraph("Địa chỉ nơi cung cấp", style_th),
            Paragraph("Hạn sử dụng", style_th),
        ]
    ]
    for r in data.get("b1_5_purchases", []):
        t15_data.append(
            [
                Paragraph(str(r["stt"]), style_td_center),
                Paragraph(r["purchase_date"], style_td_center),
                Paragraph(r["material_name"], style_td),
                Paragraph(r["quantity"], style_td_right),
                Paragraph(r["supplier_name"], style_td),
                Paragraph(r["supplier_address"], style_td),
                Paragraph(r["expiry_date"], style_td_center),
            ]
        )
    w15_pdf = [28, 72, 170, 85, 140, 180, 75]  # Tổng 750 pt
    t15_tbl = Table(t15_data, colWidths=w15_pdf, repeatRows=1)
    t15_tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F5E9")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#AAAAAA")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    elements.append(t15_tbl)

    # 4. BIỂU MẪU 1.6: SỬ DỤNG PHÂN & THUỐC (HEADER 2 TẦNG)
    elements.append(Paragraph("1.6. NHẬT KÝ: SỬ DỤNG PHÂN BÓN VÀ THUỐC BVTV, HÓA CHẤT", style_sec))
    t16_data = [
        # Dòng header 1
        [
            Paragraph("STT", style_th),
            Paragraph("Ngày, tháng, năm", style_th),
            Paragraph("Bón phân", style_th),
            "",
            Paragraph("Sử dụng thuốc BVTV", style_th),
            "",
            "",
        ],
        # Dòng header 2
        [
            "",
            "",
            Paragraph("Tên phân bón", style_th),
            Paragraph("Lượng sử dụng", style_th),
            Paragraph("Tên thuốc", style_th),
            Paragraph("Nồng độ & lượng sử dụng", style_th),
            Paragraph("Thời gian cách ly", style_th),
        ],
    ]
    for r in data.get("b1_6_usages", []):
        t16_data.append(
            [
                Paragraph(str(r["stt"]), style_td_center),
                Paragraph(r["activity_date"], style_td_center),
                Paragraph(r["fert_name"], style_td),
                Paragraph(r["fert_quantity"], style_td_right),
                Paragraph(r["pest_name"], style_td),
                Paragraph(r["pest_dose"], style_td_right),
                Paragraph(r["isolation_days"], style_td_center),
            ]
        )
    w16_pdf = [28, 72, 160, 95, 155, 145, 85]  # Tổng 740 pt
    t16_tbl = Table(t16_data, colWidths=w16_pdf, repeatRows=2)
    t16_tbl.setStyle(
        TableStyle(
            [
                ("SPAN", (0, 0), (0, 1)),  # STT
                ("SPAN", (1, 0), (1, 1)),  # Ngày
                ("SPAN", (2, 0), (3, 0)),  # Bón phân
                ("SPAN", (4, 0), (6, 0)),  # Sử dụng thuốc BVTV
                ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#E8F5E9")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#AAAAAA")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    elements.append(t16_tbl)

    # 5. BIỂU MẪU 1.7: XỬ LÝ BAO BÌ
    elements.append(Paragraph("1.7. NHẬT KÝ: THU GOM, XỬ LÝ BAO BÌ CHỨA ĐỰNG VÀ THUỐC BVTV DƯ THỪA", style_sec))
    t17_data = [
        [
            Paragraph("STT", style_th),
            Paragraph("Ngày, tháng, năm", style_th),
            Paragraph("Loại bao bì, thùng chứa, thuốc dư thừa", style_th),
            Paragraph("Nơi tồn trữ, hủy bỏ", style_th),
            Paragraph("Cách xử lý", style_th),
        ]
    ]
    for r in data.get("b1_7_waste", []):
        t17_data.append(
            [
                Paragraph(str(r["stt"]), style_td_center),
                Paragraph(r["disposal_date"], style_td_center),
                Paragraph(r["waste_type"], style_td),
                Paragraph(r["storage_location"], style_td),
                Paragraph(r["treatment_method"], style_td),
            ]
        )
    w17_pdf = [28, 72, 230, 190, 230]  # Tổng 750 pt
    t17_tbl = Table(t17_data, colWidths=w17_pdf, repeatRows=1)
    t17_tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F5E9")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#AAAAAA")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    elements.append(t17_tbl)

    # 6. BIỂU MẪU 1.8: THU HOẠCH & TIÊU THỤ
    elements.append(Paragraph("1.8. NHẬT KÝ: THU HOẠCH VÀ TIÊU THỤ SẢN PHẨM", style_sec))
    t18_data = [
        [
            Paragraph("STT", style_th),
            Paragraph("Ngày thu hoạch", style_th),
            Paragraph("Sản lượng thu hoạch", style_th),
            Paragraph("Địa điểm, cách thức sơ chế", style_th),
            Paragraph("Thời gian xuất bán", style_th),
            Paragraph("Cơ sở thu mua / tiêu thụ", style_th),
            Paragraph("Khối lượng tiêu thụ", style_th),
        ]
    ]
    for r in data.get("b1_8_harvest_sales", []):
        t18_data.append(
            [
                Paragraph(str(r["stt"]), style_td_center),
                Paragraph(r["harvest_date"], style_td_center),
                Paragraph(r["harvest_qty"], style_td_right),
                Paragraph(r["processing_location"], style_td),
                Paragraph(r["sale_date"], style_td_center),
                Paragraph(r["buyer_info"], style_td),
                Paragraph(r["sale_qty"], style_td_right),
            ]
        )
    w18_pdf = [28, 72, 100, 160, 75, 175, 90]  # Tổng 700 pt
    t18_tbl = Table(t18_data, colWidths=w18_pdf, repeatRows=1)
    t18_tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F5E9")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#AAAAAA")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    elements.append(t18_tbl)

    # 7. KHUNG KÝ TÊN
    elements.append(Spacer(1, 12))
    now = datetime.now()
    sig_date = Paragraph(f"<i>Ngày {now.day:02d} tháng {now.month:02d} năm {now.year}</i>", ParagraphStyle("SD", fontName="VN-Times-Italic", fontSize=8.5, alignment=1))
    s_tit = ParagraphStyle("ST", fontName="VN-Times-Bold", fontSize=9, alignment=1)
    s_nt = ParagraphStyle("SN", fontName="VN-Times-Italic", fontSize=8, alignment=1, textColor=colors.HexColor("#666666"))

    sig_pdf_tbl = Table(
        [
            ["", sig_date],
            [Paragraph("CHỦ NÔNG HỘ / NGƯỜI GHI CHÉP", s_tit), Paragraph("CÁN BỘ GIÁM SÁT VIETGAP", s_tit)],
            [Paragraph("(Ký và ghi rõ họ tên)", s_nt), Paragraph("(Ký, đóng dấu xác nhận)", s_nt)],
            ["", ""],
            ["", ""],
            [Paragraph(f"<b>{cover.get('farmer_name', '')}</b>", s_tit), ""],
        ],
        colWidths=[printable_width * 0.5, printable_width * 0.5],
        rowHeights=[14, 14, 12, 16, 16, 14],
    )
    sig_pdf_tbl.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
    elements.append(KeepTogether([sig_pdf_tbl]))

    doc.build(elements, canvasmaker=NumberedCanvas)


def export_vietgap_tables_pdf_bytes(
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> Tuple[io.BytesIO, str]:
    """Xuất các biểu mẫu VietGAP ra file PDF dạng buffer."""
    data = get_vietgap_full_tables_data(season_id=season_id, farmer_id=farmer_id)
    buf = io.BytesIO()
    build_vietgap_tables_pdf(data, buf)
    buf.seek(0)

    season_name = data.get("cover_info", {}).get("crop_season", "VietGAP")
    safe = re.sub(r"[^\w\-_]+", "_", season_name).strip("_")
    date_str = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"So_Tay_VietGAP_Cac_Bang_{safe}_{date_str}.pdf"

    return buf, filename


def export_vietgap_tables_pdf_file(
    output_path: str,
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> str:
    """Xuất các biểu mẫu VietGAP ra file PDF trên đĩa cứng."""
    data = get_vietgap_full_tables_data(season_id=season_id, farmer_id=farmer_id)
    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    build_vietgap_tables_pdf(data, output_path)
    logger.info("Đã xuất sổ tay VietGAP ra PDF thành công: %s", output_path)
    return output_path
# Aliases for API flexibility
get_vietgap_tables_export_data = get_vietgap_full_tables_data
build_vietgap_excel_workbook = build_vietgap_tables_workbook
build_vietgap_docx_document = build_vietgap_tables_docx
build_vietgap_pdf_document = build_vietgap_tables_pdf

export_vietgap_tables_to_excel_bytes = export_vietgap_tables_excel_bytes
export_vietgap_tables_to_excel_file = export_vietgap_tables_excel_file
export_vietgap_tables_to_excel = export_vietgap_tables_excel_file
export_vietgap_tables_to_docx_bytes = export_vietgap_tables_docx_bytes
export_vietgap_tables_to_docx_file = export_vietgap_tables_docx_file
export_vietgap_tables_to_docx = export_vietgap_tables_docx_file
export_vietgap_tables_to_pdf_bytes = export_vietgap_tables_pdf_bytes
export_vietgap_tables_to_pdf_file = export_vietgap_tables_pdf_file
export_vietgap_tables_to_pdf = export_vietgap_tables_pdf_file


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Export VietGAP Tables (Excel, Word, PDF)")
    parser.add_argument("--format", choices=["excel", "docx", "pdf", "all"], default="all", help="Định dạng xuất")
    parser.add_argument("--season-id", type=int, default=None, help="ID vụ mùa")
    parser.add_argument("--farmer-id", type=int, default=None, help="ID nông hộ")
    parser.add_argument("--prefix", type=str, default="so_tay_vietgap_cac_bang", help="Tiền tố file đầu ra")
    args = parser.parse_args()

    fmt = args.format.lower()
    if fmt in ("excel", "all"):
        f_xlsx = f"{args.prefix}.xlsx"
        export_vietgap_tables_excel_file(f_xlsx, season_id=args.season_id, farmer_id=args.farmer_id)
        print(f"✅ Đã tạo file Excel: {f_xlsx}")

    if fmt in ("docx", "all"):
        f_docx = f"{args.prefix}.docx"
        export_vietgap_tables_docx_file(f_docx, season_id=args.season_id, farmer_id=args.farmer_id)
        print(f"✅ Đã tạo file Word DOCX: {f_docx}")

    if fmt in ("pdf", "all"):
        f_pdf = f"{args.prefix}.pdf"
        export_vietgap_tables_pdf_file(f_pdf, season_id=args.season_id, farmer_id=args.farmer_id)
        print(f"✅ Đã tạo file PDF: {f_pdf}")
