"""
Module dịch vụ xuất dữ liệu nhật ký canh tác (mẫu vườn cà gai)
gộp cả 3 định dạng: Excel (.xlsx), Word (.docx), PDF (.pdf) vào 1 file duy nhất.
Cấu trúc biểu mẫu gồm:
1. Thông tin chung về vùng trồng và vụ canh tác:
   - Vụ canh tác / Crop season
   - Tên hộ sản xuất / Farmer’s name / Supervisor
   - Địa chỉ / Address
   - Mã số vùng trồng / P.U.C
   - Diện tích canh tác / Planted area
   - Giống / Variety
   - Thời điểm bắt đầu trồng
   - Thời điểm bắt đầu thu hoạch / Beginning of season
   - Sản lượng dự kiến (estimated productivity)
   - Sản lượng thực tế (actual productivity)
   - Người thu mua (Buyer)
   - Ghi chú: nội dung nào không có hoạt động thì ghi là “không”

2. Bảng ghi chép chi tiết theo ngày:
   - Ngày/tháng/năm
   - Giai đoạn sinh trưởng của cây trồng (hoặc Giai đoạn chuẩn bị trồng)
   - Công việc
   - Sinh vật gây hại phát hiện
   - Tên phân/thuốc BVTV (Bảo vệ thực vật)
   - Lượng sử dụng
   - Nơi mua
   - Thời gian cách ly
"""

import io
import os
import sys
import re
import logging
from datetime import datetime, date
from decimal import Decimal
from typing import Optional, Dict, Any, List, Tuple, Union

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

# Thư viện xử lý Excel
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Thư viện xử lý Word
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Thư viện xử lý PDF
from reportlab.lib import pagesizes, colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    KeepTogether,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from services.database import get_db_connection

logger = logging.getLogger("cultivation-diary-exporter")

# ==============================================================================
# HÀM BỔ TRỢ ĐỊNH DẠNG & FONT
# ==============================================================================

def _format_date(val: Any) -> str:
    """Định dạng ngày DD/MM/YYYY."""
    if not val:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%d/%m/%Y")
    val_str = str(val).strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", val_str)
    if m:
        return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"
    return val_str


def _format_quantity(qty: Any, unit: Optional[str] = None) -> str:
    """Định dạng số lượng."""
    if qty is None or qty == "":
        return "-"
    try:
        val_f = float(qty)
        fmt_val = f"{int(val_f):,}" if val_f.is_integer() else f"{val_f:,.2f}"
        fmt_val = fmt_val.replace(",", "X").replace(".", ",").replace("X", ".")
        u = (unit or "").strip()
        return f"{fmt_val} {u}".strip() if u else fmt_val
    except (ValueError, TypeError):
        return str(qty)


_FONTS_REGISTERED = False


def _init_fonts():
    """Khởi tạo font tiếng Việt Unicode cho ReportLab."""
    global _FONTS_REGISTERED
    if _FONTS_REGISTERED:
        return

    candidates = [
        (
            "/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman.ttf",
            "/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman_Bold.ttf",
            "/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman_Italic.ttf",
            "/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman_Bold_Italic.ttf",
        ),
        (
            "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSerif-Italic.ttf",
            "/usr/share/fonts/truetype/liberation/LiberationSerif-BoldItalic.ttf",
        ),
        (
            "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Italic.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSerif-BoldItalic.ttf",
        ),
    ]

    for reg, bld, itl, blditl in candidates:
        if os.path.exists(reg):
            try:
                pdfmetrics.registerFont(TTFont("VN-Times", reg))
                pdfmetrics.registerFont(TTFont("VN-Times-Bold", bld if os.path.exists(bld) else reg))
                pdfmetrics.registerFont(TTFont("VN-Times-Italic", itl if os.path.exists(itl) else reg))
                pdfmetrics.registerFont(TTFont("VN-Times-BoldItalic", blditl if os.path.exists(blditl) else reg))
                pdfmetrics.registerFontFamily(
                    "VN-Times",
                    normal="VN-Times",
                    bold="VN-Times-Bold",
                    italic="VN-Times-Italic",
                    boldItalic="VN-Times-BoldItalic",
                )
                _FONTS_REGISTERED = True
                return
            except Exception as e:
                logger.warning("Lỗi đăng ký font PDF: %s", e)

    _FONTS_REGISTERED = True


class NumberedCanvas(canvas.Canvas):
    """Canvas 2-pass vẽ footer và số trang Trang X / Y."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_footer(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_footer(self, page_count):
        self.saveState()
        _init_fonts()
        self.setFont("VN-Times", 8)
        self.setFillColor(colors.HexColor("#555555"))

        page_width, _ = pagesizes.landscape(pagesizes.A4)
        margin = 36
        self.setStrokeColor(colors.HexColor("#CCCCCC"))
        self.setLineWidth(0.5)
        self.line(margin, 25, page_width - margin, 25)

        self.drawString(margin, 15, "Hệ thống quản lý Nông ký điện tử VietGAP — Chuẩn quốc gia")
        page_str = f"Trang {self._pageNumber} / {page_count}"
        self.drawRightString(page_width - margin, 15, page_str)
        self.restoreState()


# ==============================================================================
# TRUY VẤN DỮ LIỆU TỪ CSDL
# ==============================================================================

def get_available_seasons(conn=None, farmer_id: Optional[Any] = None) -> List[Dict[str, Any]]:
    """Lấy danh sách các mùa vụ canh tác khả dụng để xuất."""
    should_close = False
    if conn is None:
        conn = get_db_connection()
        should_close = True

    try:
        with conn.cursor() as cur:
            sql = """
                SELECT 
                    cs.season_id,
                    cs.season_name,
                    cs.crop_name,
                    cs.variety,
                    cs.start_date,
                    p.plot_id,
                    p.name AS plot_name,
                    f.farm_id,
                    f.name AS farm_name,
                    f.farm_code,
                    u.full_name AS farmer_name
                FROM crop_seasons cs
                JOIN plots p ON cs.plot_id = p.plot_id
                JOIN farm f ON p.farm_id = f.farm_id
                LEFT JOIN users u ON f.owner_user_id = u.user_id
            """
            params = []
            if farmer_id:
                sql += " WHERE f.owner_user_id = %s "
                params.append(farmer_id)
            sql += " ORDER BY cs.season_id DESC;"

            cur.execute(sql, tuple(params))
            rows = cur.fetchall()
            return [
                {
                    "season_id": r["season_id"],
                    "season_name": r["season_name"],
                    "crop_name": r.get("crop_name") or "",
                    "variety": r.get("variety") or "",
                    "start_date": _format_date(r.get("start_date")),
                    "plot_name": r.get("plot_name") or "",
                    "farm_name": r.get("farm_name") or "",
                    "farm_code": r.get("farm_code") or "",
                    "farmer_name": r.get("farmer_name") or "",
                }
                for r in rows
            ]
    finally:
        if should_close:
            conn.close()


def get_cultivation_diary_export_data(
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> Dict[str, Any]:
    """
    Truy vấn CSDL gom dữ liệu đầy đủ cho 2 phần của báo cáo:
    1. Thông tin chung về vùng trồng và vụ canh tác (11 trường)
    2. Chi tiết ghi chép hoạt động theo ngày (8 cột)
    """
    conn = get_db_connection()
    try:
        with conn.cursor() as cur:
            # 1. Tìm season_id phù hợp
            if not season_id and farmer_id:
                cur.execute(
                    """
                    SELECT cs.season_id 
                    FROM crop_seasons cs
                    JOIN plots p ON cs.plot_id = p.plot_id
                    JOIN farm f ON p.farm_id = f.farm_id
                    WHERE f.owner_user_id = %s
                    ORDER BY cs.season_id DESC LIMIT 1;
                    """,
                    (farmer_id,),
                )
                res = cur.fetchone()
                if res:
                    season_id = res["season_id"]

            if not season_id:
                cur.execute("SELECT season_id FROM crop_seasons ORDER BY season_id DESC LIMIT 1;")
                res = cur.fetchone()
                if res:
                    season_id = res["season_id"]

            if not season_id:
                raise ValueError("Không tìm thấy vụ canh tác nào trong CSDL để xuất.")

            # 2. Thông tin chung
            cur.execute(
                """
                SELECT 
                    cs.season_id,
                    cs.season_name,
                    cs.crop_name,
                    cs.variety AS season_variety,
                    cs.start_date,
                    cs.expected_harvest_date,
                    cs.expected_yield,
                    cs.actual_yield,
                    cs.yield_unit,
                    p.plot_id,
                    p.name AS plot_name,
                    p.plot_code,
                    p.variety AS plot_variety,
                    p.area_size,
                    p.area_unit,
                    f.farm_id,
                    f.name AS farm_name,
                    f.farm_code,
                    f.address AS farm_address,
                    u.user_id,
                    u.full_name AS farmer_name,
                    u.username AS farmer_username,
                    u.phone AS farmer_phone
                FROM crop_seasons cs
                JOIN plots p ON cs.plot_id = p.plot_id
                JOIN farm f ON p.farm_id = f.farm_id
                LEFT JOIN users u ON f.owner_user_id = u.user_id
                WHERE cs.season_id = %s
                LIMIT 1;
                """,
                (season_id,),
            )
            season_info = cur.fetchone()
            if not season_info:
                raise ValueError(f"Không tìm thấy thông tin cho season_id={season_id}")

            # Người thu mua
            cur.execute(
                """
                SELECT DISTINCT sr.buyer_name, sr.buyer_contact
                FROM sales_records sr
                JOIN harvest_records hr ON sr.harvest_id = hr.harvest_id
                WHERE hr.season_id = %s AND sr.buyer_name IS NOT NULL AND TRIM(sr.buyer_name) != '';
                """,
                (season_id,),
            )
            buyer_rows = cur.fetchall()
            buyer_names = []
            for b in buyer_rows:
                b_name = b["buyer_name"].strip()
                if b.get("buyer_contact"):
                    b_name += f" ({b['buyer_contact']})"
                if b_name not in buyer_names:
                    buyer_names.append(b_name)
            buyer_str = ", ".join(buyer_names) if buyer_names else "Chưa có thông tin thương lái / Hợp tác xã thu mua"

            # Thống kê thu hoạch
            cur.execute(
                """
                SELECT 
                    MIN(ca.activity_date) AS min_harvest_date,
                    SUM(hr.quantity) AS total_harvest_qty,
                    hr.unit AS harvest_unit
                FROM harvest_records hr
                JOIN cultivation_activities ca ON hr.activity_id = ca.activity_id
                WHERE hr.season_id = %s
                GROUP BY hr.unit;
                """,
                (season_id,),
            )
            harvest_stats = cur.fetchall()
            first_harvest_date = None
            total_harvest_amount = Decimal(0)
            harvest_unit = season_info.get("yield_unit") or "kg"
            for h in harvest_stats:
                if h.get("min_harvest_date"):
                    if not first_harvest_date or h["min_harvest_date"] < first_harvest_date:
                        first_harvest_date = h["min_harvest_date"]
                if h.get("total_harvest_qty"):
                    total_harvest_amount += h["total_harvest_qty"]
                if h.get("harvest_unit"):
                    harvest_unit = h["harvest_unit"]

            harvest_start_display = ""
            if first_harvest_date:
                harvest_start_display = _format_date(first_harvest_date)
            elif season_info.get("expected_harvest_date"):
                harvest_start_display = f"{_format_date(season_info['expected_harvest_date'])} (Dự kiến)"
            else:
                harvest_start_display = "Chưa ghi nhận"

            if season_info.get("actual_yield") is not None:
                actual_yield_str = _format_quantity(season_info["actual_yield"], season_info.get("yield_unit") or "kg")
            elif total_harvest_amount > 0:
                actual_yield_str = _format_quantity(total_harvest_amount, harvest_unit)
            else:
                actual_yield_str = "Chưa thu hoạch"

            expected_yield_str = _format_quantity(
                season_info.get("expected_yield"), season_info.get("yield_unit") or "kg"
            )
            area_str = _format_quantity(season_info.get("area_size"), season_info.get("area_unit") or "m²")

            variety_str = (
                season_info.get("season_variety")
                or season_info.get("plot_variety")
                or season_info.get("crop_name")
                or "-"
            )
            if season_info.get("crop_name") and season_info.get("crop_name") not in variety_str:
                variety_str = f"{season_info['crop_name']} - {variety_str}"

            farmer_display = (
                season_info.get("farmer_name")
                or season_info.get("farm_name")
                or season_info.get("farmer_username")
                or "Nông hộ sản xuất"
            )

            general_info = {
                "crop_season": season_info.get("season_name") or f"Vụ mùa #{season_id}",
                "farmer_name": farmer_display,
                "address": season_info.get("farm_address") or "Chưa cập nhật",
                "farm_code": season_info.get("farm_code") or season_info.get("plot_code") or "VN-GAP-001",
                "planted_area": area_str,
                "variety": variety_str,
                "start_date": _format_date(season_info.get("start_date")),
                "harvest_start_date": harvest_start_display,
                "expected_yield": expected_yield_str,
                "actual_yield": actual_yield_str,
                "buyer": buyer_str,
                "farm_name": season_info.get("farm_name") or "",
                "plot_name": season_info.get("plot_name") or "",
            }

            # Map nhà cung cấp
            cur.execute(
                """
                SELECT ip.pesticide_id, ip.fertilizer_id, s.name AS supplier_name
                FROM input_purchases ip
                JOIN suppliers s ON ip.supplier_id = s.supplier_id
                WHERE ip.season_id = %s
                ORDER BY ip.purchase_date DESC;
                """,
                (season_id,),
            )
            purchased_suppliers = cur.fetchall()
            pest_supplier_map = {}
            fert_supplier_map = {}
            for ps in purchased_suppliers:
                if ps.get("pesticide_id") and ps["pesticide_id"] not in pest_supplier_map:
                    pest_supplier_map[ps["pesticide_id"]] = ps["supplier_name"]
                if ps.get("fertilizer_id") and ps["fertilizer_id"] not in fert_supplier_map:
                    fert_supplier_map[ps["fertilizer_id"]] = ps["supplier_name"]

            # Chi tiết hoạt động
            cur.execute(
                """
                SELECT 
                    ca.activity_id,
                    ca.activity_date,
                    ca.activity_time,
                    ca.activity_type,
                    ca.growth_stage,
                    ca.description,
                    ca.weather_condition,
                    ca.performed_by,
                    vr.stt_text,

                    pa.application_id AS pest_app_id,
                    pa.pesticide_id,
                    pa.dose_quantity,
                    pa.dose_unit,
                    pa.concentration_value,
                    pa.concentration_unit,
                    pa.application_method AS pest_method,
                    pa.target_pest,
                    pa.isolation_days,
                    p.name AS pesticide_name,
                    p.min_isolation_days,
                    s_pa.name AS pesticide_supplier,

                    fa.application_id AS fert_app_id,
                    fa.fertilizer_id,
                    fa.quantity AS fert_quantity,
                    fa.unit AS fert_unit,
                    fa.application_method AS fert_method,
                    f.name AS fertilizer_name,
                    s_fa.name AS fertilizer_supplier,

                    po.observation_id,
                    po.pest_name,
                    po.pest_type,
                    po.severity,
                    po.action_taken,

                    ir.irrigation_id,
                    ir.water_source,
                    ir.method AS irrigation_method,
                    ir.duration_minutes,
                    ir.quantity AS irrigation_quantity,

                    hr.harvest_id,
                    hr.quantity AS harvest_quantity,
                    hr.unit AS harvest_unit,
                    hr.processing_location

                FROM cultivation_activities ca
                LEFT JOIN activity_voice_records vr ON ca.activity_id = vr.activity_id

                LEFT JOIN pesticide_applications pa ON ca.activity_id = pa.activity_id
                LEFT JOIN pesticides p ON pa.pesticide_id = p.pesticide_id
                LEFT JOIN suppliers s_pa ON pa.supplier_id = s_pa.supplier_id

                LEFT JOIN fertilizer_applications fa ON ca.activity_id = fa.activity_id
                LEFT JOIN fertilizers f ON fa.fertilizer_id = f.fertilizer_id
                LEFT JOIN suppliers s_fa ON fa.supplier_id = s_fa.supplier_id

                LEFT JOIN pest_observations po ON ca.activity_id = po.activity_id
                LEFT JOIN irrigation_records ir ON ca.activity_id = ir.activity_id
                LEFT JOIN harvest_records hr ON ca.activity_id = hr.activity_id

                WHERE ca.season_id = %s
                ORDER BY ca.activity_date ASC, ca.activity_time ASC, ca.activity_id ASC;
                """,
                (season_id,),
            )
            raw_activities = cur.fetchall()

            daily_records: List[Dict[str, Any]] = []
            for idx, r in enumerate(raw_activities, start=1):
                act_type = r.get("activity_type") or "OTHER"
                date_str = _format_date(r.get("activity_date"))
                growth_stage = r.get("growth_stage") or "-"

                desc = (r.get("description") or "").strip()
                if not desc:
                    if act_type == "PESTICIDE":
                        desc = f"Phun thuốc BVTV ({r.get('pesticide_name') or 'phòng trừ sâu bệnh'})"
                    elif act_type == "FERTILIZER":
                        desc = f"Bón phân ({r.get('fertilizer_name') or 'dinh dưỡng'})"
                    elif act_type == "PEST_OBSERVATION":
                        desc = f"Kiểm tra sâu bệnh ({r.get('pest_name') or 'định kỳ'})"
                    elif act_type == "IRRIGATION":
                        desc = f"Tưới nước ({r.get('irrigation_method') or 'vườn'})"
                    elif act_type == "HARVEST":
                        desc = "Thu hoạch nông sản"
                    else:
                        desc = "Hoạt động chăm sóc canh tác"

                pest_detected = "-"
                if r.get("pest_name"):
                    pest_detected = r["pest_name"].strip()
                elif r.get("target_pest"):
                    pest_detected = r["target_pest"].strip()
                else:
                    stt = f"{desc} {r.get('stt_text') or ''}".lower()
                    m_p = re.search(r"(rầy\s+\w+|sâu\s+\w+|bệnh\s+\w+|nấm\s+\w+|bọ\s+\w+|nhện\s+\w+)", stt)
                    if m_p and not any(k in stt for k in ["không có", "không phát hiện", "sạch", "phòng ngừa"]):
                        pest_detected = m_p.group(1).title()
                    elif act_type in ("PESTICIDE", "PEST_OBSERVATION"):
                        pest_detected = "Không phát hiện"

                chemical_name = "-"
                if r.get("pesticide_name"):
                    chemical_name = r["pesticide_name"].strip()
                elif r.get("fertilizer_name"):
                    chemical_name = r["fertilizer_name"].strip()

                quantity_used = "-"
                if act_type == "PESTICIDE" or r.get("pest_app_id"):
                    dose_q = r.get("dose_quantity")
                    dose_u = r.get("dose_unit") or "lít"
                    conc_v = r.get("concentration_value")
                    conc_u = r.get("concentration_unit") or "%"
                    if dose_q is not None:
                        quantity_used = _format_quantity(dose_q, dose_u)
                        if conc_v is not None:
                            quantity_used += f" (Nồng độ: {_format_quantity(conc_v, conc_u)})"
                elif act_type == "FERTILIZER" or r.get("fert_app_id"):
                    f_q = r.get("fert_quantity")
                    f_u = r.get("fert_unit") or "kg"
                    if f_q is not None:
                        quantity_used = _format_quantity(f_q, f_u)
                elif act_type == "IRRIGATION" or r.get("irrigation_id"):
                    ir_q = r.get("irrigation_quantity")
                    ir_dur = r.get("duration_minutes")
                    if ir_q is not None:
                        quantity_used = _format_quantity(ir_q, "m³")
                    elif ir_dur is not None:
                        quantity_used = f"{ir_dur} phút"
                elif act_type == "HARVEST" or r.get("harvest_id"):
                    h_q = r.get("harvest_quantity")
                    h_u = r.get("harvest_unit") or "kg"
                    if h_q is not None:
                        quantity_used = _format_quantity(h_q, h_u)

                where_bought = "-"
                if r.get("pesticide_supplier"):
                    where_bought = r["pesticide_supplier"].strip()
                elif r.get("fertilizer_supplier"):
                    where_bought = r["fertilizer_supplier"].strip()
                elif r.get("pesticide_id") and r["pesticide_id"] in pest_supplier_map:
                    where_bought = pest_supplier_map[r["pesticide_id"]]
                elif r.get("fertilizer_id") and r["fertilizer_id"] in fert_supplier_map:
                    where_bought = fert_supplier_map[r["fertilizer_id"]]

                isolation_period = "-"
                if r.get("isolation_days") is not None:
                    isolation_period = f"{r['isolation_days']} ngày"
                elif r.get("min_isolation_days") is not None:
                    isolation_period = f"{r['min_isolation_days']} ngày"
                elif act_type == "PESTICIDE":
                    isolation_period = "7 ngày"

                daily_records.append(
                    {
                        "stt": idx,
                        "date": date_str,
                        "growth_stage": growth_stage,
                        "work_description": desc,
                        "pest_detected": pest_detected,
                        "chemical_name": chemical_name,
                        "quantity_used": quantity_used,
                        "where_bought": where_bought,
                        "isolation_period": isolation_period,
                    }
                )

        return {
            "season_id": season_id,
            "general_info": general_info,
            "daily_records": daily_records,
            "exported_at": datetime.now().strftime("%d/%m/%Y %H:%M"),
        }
    finally:
        conn.close()


# ==============================================================================
# 1. XUẤT EXCEL (.XLSX)
# ==============================================================================

def build_cultivation_diary_workbook(data: Dict[str, Any]) -> openpyxl.Workbook:
    """Tạo workbook Excel định dạng chuẩn VietGAP."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nhật Ký Canh Tác VietGAP"
    ws.views.sheetView[0].showGridLines = True

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    COLOR_PRIMARY_DARK = "1B5E20"
    COLOR_PRIMARY_LIGHT = "2E7D32"
    COLOR_ACCENT_BG = "E8F5E9"
    COLOR_ZEBRA_ROW = "F9FBF7"
    COLOR_BORDER = "CCCCCC"

    font_main_title = Font(name="Arial", size=16, bold=True, color=COLOR_PRIMARY_DARK)
    font_sub_title = Font(name="Arial", size=10, italic=True, color="555555")
    font_section_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_field_label = Font(name="Arial", size=10, bold=True, color=COLOR_PRIMARY_DARK)
    font_field_value = Font(name="Arial", size=10, color="111111")
    font_table_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Arial", size=9, color="212121")
    font_data_bold = Font(name="Arial", size=9, bold=True, color="1B5E20")
    font_footer = Font(name="Arial", size=9, italic=True, color="666666")
    font_sign_title = Font(name="Arial", size=10, bold=True, color="111111")
    font_sign_note = Font(name="Arial", size=9, italic=True, color="777777")

    fill_section_header = PatternFill(start_color=COLOR_PRIMARY_LIGHT, end_color=COLOR_PRIMARY_LIGHT, fill_type="solid")
    fill_field_label = PatternFill(start_color=COLOR_ACCENT_BG, end_color=COLOR_ACCENT_BG, fill_type="solid")
    fill_table_header = PatternFill(start_color=COLOR_PRIMARY_DARK, end_color=COLOR_PRIMARY_DARK, fill_type="solid")
    fill_zebra = PatternFill(start_color=COLOR_ZEBRA_ROW, end_color=COLOR_ZEBRA_ROW, fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border_side = Side(border_style="thin", color=COLOR_BORDER)
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    header_border = Border(
        left=Side(border_style="thin", color="4CAF50"),
        right=Side(border_style="thin", color="4CAF50"),
        top=Side(border_style="medium", color=COLOR_PRIMARY_DARK),
        bottom=Side(border_style="medium", color=COLOR_PRIMARY_DARK),
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    gen = data.get("general_info", {})

    ws.merge_cells("A1:I1")
    farm_tag_text = (gen.get("farm_name") or "HỆ THỐNG QUẢN LÝ NÔNG KÝ ĐIỆN TỬ VIETGAP").upper()
    ws["A1"].value = f"TRANG TRẠI: {farm_tag_text}"
    ws["A1"].font = Font(name="Arial", size=10, bold=True, color="388E3C")
    ws["A1"].alignment = align_center
    ws.row_dimensions[1].height = 20

    ws.merge_cells("A2:I2")
    ws["A2"].value = "SỔ NHẬT KÝ CANH TÁC THEO TIÊU CHUẨN VIETGAP"
    ws["A2"].font = font_main_title
    ws["A2"].alignment = align_center
    ws.row_dimensions[2].height = 28

    ws.merge_cells("A3:I3")
    ws["A3"].value = "CULTIVATION & PRODUCTION LOGBOOK (GOOD AGRICULTURAL PRACTICES)"
    ws["A3"].font = font_sub_title
    ws["A3"].alignment = align_center
    ws.row_dimensions[3].height = 18

    # PHẦN 1
    row_sec1 = 5
    ws.merge_cells(f"A{row_sec1}:I{row_sec1}")
    ws[f"A{row_sec1}"].value = "1. PHẦN THÔNG TIN CHUNG VỀ VÙNG TRỒNG VÀ VỤ CANH TÁC"
    ws[f"A{row_sec1}"].font = font_section_header
    ws[f"A{row_sec1}"].fill = fill_section_header
    ws[f"A{row_sec1}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row_sec1].height = 24

    for col in range(1, 10):
        ws.cell(row=row_sec1, column=col).border = thin_border

    general_fields = [
        ("Vụ canh tác / Crop season:", gen.get("crop_season", "-"), "Thời điểm bắt đầu trồng:", gen.get("start_date", "-")),
        ("Tên hộ sản xuất / Farmer's name:", gen.get("farmer_name", "-"), "Thời điểm bắt đầu thu hoạch / Season start:", gen.get("harvest_start_date", "-")),
        ("Địa chỉ / Address:", gen.get("address", "-"), "Sản lượng dự kiến (estimated productivity):", gen.get("expected_yield", "-")),
        ("Mã số vùng trồng / P.U.C:", gen.get("farm_code", "-"), "Sản lượng thực tế (actual productivity):", gen.get("actual_yield", "-")),
        ("Diện tích canh tác / Planted area:", gen.get("planted_area", "-"), "Người thu mua (Buyer):", gen.get("buyer", "-")),
        ("Giống / Variety:", gen.get("variety", "-"), "Lô canh tác / Plot:", gen.get("plot_name", "-")),
    ]

    current_r = row_sec1 + 1
    for left_lbl, left_val, right_lbl, right_val in general_fields:
        ws.row_dimensions[current_r].height = 22
        ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=2)
        c = ws.cell(row=current_r, column=1, value=left_lbl)
        c.font = font_field_label
        c.fill = fill_field_label
        c.alignment = align_left

        ws.merge_cells(start_row=current_r, start_column=3, end_row=current_r, end_column=4)
        c = ws.cell(row=current_r, column=3, value=left_val)
        c.font = font_field_value
        c.fill = fill_white
        c.alignment = align_left

        ws.merge_cells(start_row=current_r, start_column=5, end_row=current_r, end_column=6)
        c = ws.cell(row=current_r, column=5, value=right_lbl)
        c.font = font_field_label
        c.fill = fill_field_label
        c.alignment = align_left

        ws.merge_cells(start_row=current_r, start_column=7, end_row=current_r, end_column=9)
        c = ws.cell(row=current_r, column=7, value=right_val)
        c.font = font_field_value
        c.fill = fill_white
        c.alignment = align_left

        for col in range(1, 10):
            ws.cell(row=current_r, column=col).border = thin_border
        current_r += 1

    # PHẦN 2
    current_r += 1
    row_sec2 = current_r
    ws.merge_cells(f"A{row_sec2}:I{row_sec2}")
    ws[f"A{row_sec2}"].value = "2. PHẦN THÔNG TIN GHI CHÉP CHI TIẾT THEO NGÀY (DAILY CULTIVATION LOG)"
    ws[f"A{row_sec2}"].font = font_section_header
    ws[f"A{row_sec2}"].fill = fill_section_header
    ws[f"A{row_sec2}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row_sec2].height = 24

    for col in range(1, 10):
        ws.cell(row=row_sec2, column=col).border = thin_border

    current_r += 1
    row_table_header = current_r
    ws.row_dimensions[row_table_header].height = 36

    table_headers = [
        ("STT", 6, align_center),
        ("Ngày/tháng/năm", 14, align_center),
        ("Giai đoạn sinh trưởng của cây trồng\n(hoặc Giai đoạn chuẩn bị trồng)", 26, align_center),
        ("Công việc", 32, align_center),
        ("Sinh vật gây hại phát hiện", 22, align_center),
        ("Tên phân/thuốc BVTV\n(Bảo vệ thực vật)", 24, align_center),
        ("Lượng sử dụng", 18, align_center),
        ("Nơi mua", 22, align_center),
        ("Thời gian cách ly", 16, align_center),
    ]

    for col_idx, (h_title, col_width, h_align) in enumerate(table_headers, start=1):
        cell_h = ws.cell(row=row_table_header, column=col_idx, value=h_title)
        cell_h.font = font_table_header
        cell_h.fill = fill_table_header
        cell_h.alignment = h_align
        cell_h.border = header_border

    daily_records = data.get("daily_records", [])
    current_r += 1

    for r_data in daily_records:
        is_even = (r_data["stt"] % 2 == 0)
        row_fill = fill_zebra if is_even else fill_white

        desc_len = len(str(r_data.get("work_description", "")))
        ws.row_dimensions[current_r].height = 42 if desc_len > 80 else (30 if desc_len > 40 else 22)

        ws.cell(row=current_r, column=1, value=r_data["stt"]).alignment = align_center
        ws.cell(row=current_r, column=2, value=r_data["date"]).alignment = align_center
        ws.cell(row=current_r, column=3, value=r_data["growth_stage"]).alignment = align_left
        ws.cell(row=current_r, column=4, value=r_data["work_description"]).alignment = align_left

        c5 = ws.cell(row=current_r, column=5, value=r_data["pest_detected"])
        c5.alignment = align_left
        c5.font = Font(name="Arial", size=9, bold=True, color="D32F2F") if r_data["pest_detected"] not in ("-", "Không phát hiện") else font_data

        c6 = ws.cell(row=current_r, column=6, value=r_data["chemical_name"])
        c6.alignment = align_left
        c6.font = font_data_bold if r_data["chemical_name"] != "-" else font_data

        ws.cell(row=current_r, column=7, value=r_data["quantity_used"]).alignment = align_right
        ws.cell(row=current_r, column=8, value=r_data["where_bought"]).alignment = align_left

        c9 = ws.cell(row=current_r, column=9, value=r_data["isolation_period"])
        c9.alignment = align_center
        c9.font = Font(name="Arial", size=9, bold=True, color="E65100") if r_data["isolation_period"] != "-" else font_data

        for col in range(1, 10):
            c = ws.cell(row=current_r, column=col)
            if col not in (5, 6, 9):
                c.font = font_data
            c.fill = row_fill
            c.border = thin_border
        current_r += 1

    # Footer & Chữ ký
    ws.row_dimensions[current_r].height = 22
    ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=3)
    c_tot = ws.cell(row=current_r, column=1, value=f"Tổng số lượt ghi chép: {len(daily_records)} lượt")
    c_tot.font = Font(name="Arial", size=9, bold=True, italic=True, color=COLOR_PRIMARY_DARK)
    c_tot.fill = fill_field_label
    c_tot.alignment = align_left

    ws.merge_cells(start_row=current_r, start_column=4, end_row=current_r, end_column=9)
    c_time = ws.cell(row=current_r, column=4, value=f"Thời gian xuất báo cáo: {data.get('exported_at', '')}")
    c_time.font = font_footer
    c_time.fill = fill_field_label
    c_time.alignment = align_right

    for col in range(1, 10):
        ws.cell(row=current_r, column=col).border = thin_border

    current_r += 2
    today_dt = datetime.now()
    ws.merge_cells(start_row=current_r, start_column=6, end_row=current_r, end_column=9)
    ws.cell(row=current_r, column=6, value=f"Ngày {today_dt.day:02d} tháng {today_dt.month:02d} năm {today_dt.year}").font = font_sign_note
    ws.cell(row=current_r, column=6).alignment = align_center

    current_r += 1
    ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=4)
    ws.cell(row=current_r, column=1, value="NGƯỜI GHI CHÉP / CHỦ HỘ SẢN XUẤT").font = font_sign_title
    ws.cell(row=current_r, column=1).alignment = align_center

    ws.merge_cells(start_row=current_r, start_column=6, end_row=current_r, end_column=9)
    ws.cell(row=current_r, column=6, value="KỸ THUẬT VIÊN / CÁN BỘ GIÁM SÁT VIETGAP").font = font_sign_title
    ws.cell(row=current_r, column=6).alignment = align_center

    current_r += 1
    ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=4)
    ws.cell(row=current_r, column=1, value="(Ký và ghi rõ họ tên)").font = font_sign_note
    ws.cell(row=current_r, column=1).alignment = align_center

    ws.merge_cells(start_row=current_r, start_column=6, end_row=current_r, end_column=9)
    ws.cell(row=current_r, column=6, value="(Ký, đóng dấu hoặc xác nhận)").font = font_sign_note
    ws.cell(row=current_r, column=6).alignment = align_center

    current_r += 4
    ws.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=4)
    ws.cell(row=current_r, column=1, value=gen.get("farmer_name", "")).font = font_sign_title
    ws.cell(row=current_r, column=1).alignment = align_center

    for col_idx, (_, col_w, _) in enumerate(table_headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = col_w

    return wb


def export_cultivation_diary_to_bytes(
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> Tuple[io.BytesIO, str]:
    """Xuất nhật ký canh tác ra buffer Excel (.xlsx)."""
    data = get_cultivation_diary_export_data(season_id=season_id, farmer_id=farmer_id)
    wb = build_cultivation_diary_workbook(data)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    season_name = data.get("general_info", {}).get("crop_season", "Vu_Mua")
    safe = re.sub(r"[^\w\-_]+", "_", season_name).strip("_")
    date_str = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"Nhat_Ky_Canh_Tac_VietGAP_{safe}_{date_str}.xlsx"
    return buf, filename


def export_cultivation_diary_to_file(
    output_path: str,
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> str:
    """Xuất nhật ký canh tác ra file Excel trên đĩa cứng."""
    data = get_cultivation_diary_export_data(season_id=season_id, farmer_id=farmer_id)
    wb = build_cultivation_diary_workbook(data)
    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(output_path)
    return output_path


# ==============================================================================
# 2. XUẤT WORD (.DOCX)
# ==============================================================================

def build_cultivation_diary_docx(data: Dict[str, Any]) -> docx.Document:
    """Khởi tạo Document Word (.docx) theo mẫu cà gai."""
    doc = docx.Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.27)
    section.right_margin = Cm(1.27)

    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(11)
    gen = data.get("general_info", {})

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("SỔ NHẬT KÝ CANH TÁC/Farm diary")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(15)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)

    meta_items = [
        ("Vụ canh tác/Crop season: ", gen.get("crop_season", "-")),
        ("Tên hộ sản xuất/Farmer’s name/Supervisor: ", gen.get("farmer_name", "-")),
        ("Địa chỉ/Address: ", gen.get("address", "-")),
        ("Mã số vùng trồng/P.U.C: ", gen.get("farm_code", "-")),
        ("Diện tích canh tác/Planted area: ", gen.get("planted_area", "-")),
        ("Giống/Variety: ", gen.get("variety", "-")),
        ("Thời điểm bắt đầu trồng: ", gen.get("start_date", "-")),
        ("Thời điểm bắt đầu thu hoạch/ Beginning of season: ", gen.get("harvest_start_date", "-")),
        ("Sản lượng dự kiến (estimated productivity): ", gen.get("expected_yield", "-")),
        ("Sản lượng thực tế (actual productivity): ", gen.get("actual_yield", "-")),
        ("Người thu mua (Buyer): ", gen.get("buyer", "-")),
    ]

    for label, val in meta_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2.5)
        r1 = p.add_run(label)
        r1.font.bold = True
        r2 = p.add_run(str(val))
        r2.font.bold = False

    p_note = doc.add_paragraph()
    p_note.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_n = p_note.add_run("Ghi chú: nội dung nào không có hoạt động thì ghi là “không”")
    r_n.font.italic = True
    r_n.font.size = Pt(10)

    headers = [
        "Ngày/tháng/năm",
        "Giai đoạn sinh trưởng\ncủa cây trồng",
        "Công việc",
        "Sinh vật gây hại\nphát hiện",
        "Tên phân/thuốc BVTV",
        "Lượng sử dụng",
        "Nơi mua",
        "Thời gian cách ly",
    ]
    col_widths = [Cm(2.8), Cm(3.6), Cm(5.2), Cm(3.2), Cm(3.5), Cm(2.8), Cm(3.5), Cm(2.5)]

    daily_records = data.get("daily_records", [])
    table = doc.add_table(rows=len(daily_records) + 1, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(OxmlElement('w:tblHeader'))

    border_spec = {"sz": 4, "val": "single", "color": "999999"}
    cell_borders = {"top": border_spec, "bottom": border_spec, "left": border_spec, "right": border_spec}

    for col_idx, (h_text, width) in enumerate(zip(headers, col_widths)):
        cell = table.cell(0, col_idx)
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8F5E9"/>')
        cell._tc.get_or_add_tcPr().append(shd)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)

    def _clean_val(v: Any) -> str:
        s = str(v).strip() if v is not None else ""
        return "Không" if s in ("", "-") else s

    for row_idx, r_data in enumerate(daily_records, start=1):
        row_vals = [
            r_data.get("date", ""),
            _clean_val(r_data.get("growth_stage")),
            _clean_val(r_data.get("work_description")),
            _clean_val(r_data.get("pest_detected")),
            _clean_val(r_data.get("chemical_name")),
            _clean_val(r_data.get("quantity_used")),
            _clean_val(r_data.get("where_bought")),
            _clean_val(r_data.get("isolation_period")),
        ]

        trPr = table.rows[row_idx]._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        row_bg = "F9FBF7" if (row_idx % 2 == 0) else "FFFFFF"

        for col_idx, (val_text, width) in enumerate(zip(row_vals, col_widths)):
            cell = table.cell(row_idx, col_idx)
            cell.width = width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{row_bg}"/>')
            cell._tc.get_or_add_tcPr().append(shd)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in (0, 7) else (WD_ALIGN_PARAGRAPH.RIGHT if col_idx == 5 else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(val_text)
            r.font.size = Pt(9.5)

    # Chữ ký
    now_dt = datetime.now()
    p_d = doc.add_paragraph()
    p_d.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_d.paragraph_format.space_before = Pt(14)
    r = p_d.add_run(f"Ngày {now_dt.day:02d} tháng {now_dt.month:02d} năm {now_dt.year}")
    r.font.italic = True

    sig_t = doc.add_table(rows=3, cols=2)
    sig_t.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_t.autofit = False
    for row in sig_t.rows:
        row.cells[0].width = Cm(13.5)
        row.cells[1].width = Cm(13.5)

    sig_t.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig_t.cell(0, 0).paragraphs[0].add_run("NGƯỜI GHI CHÉP / CHỦ HỘ SẢN XUẤT").font.bold = True
    sig_t.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig_t.cell(0, 1).paragraphs[0].add_run("KỸ THUẬT VIÊN / CÁN BỘ GIÁM SÁT VIETGAP").font.bold = True

    sig_t.cell(1, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig_t.cell(1, 0).paragraphs[0].add_run("(Ký và ghi rõ họ tên)").font.italic = True
    sig_t.cell(1, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig_t.cell(1, 1).paragraphs[0].add_run("(Ký, đóng dấu hoặc xác nhận)").font.italic = True

    sig_t.cell(2, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig_t.cell(2, 0).paragraphs[0].paragraph_format.space_before = Pt(40)
    sig_t.cell(2, 0).paragraphs[0].add_run(gen.get("farmer_name", "")).font.bold = True

    return doc


def export_cultivation_diary_to_docx_bytes(
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> Tuple[io.BytesIO, str]:
    """Xuất nhật ký canh tác ra buffer Word (.docx)."""
    data = get_cultivation_diary_export_data(season_id=season_id, farmer_id=farmer_id)
    doc = build_cultivation_diary_docx(data)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    season_name = data.get("general_info", {}).get("crop_season", "Vu_Canh_Tac")
    safe = re.sub(r"[^\w\-_]+", "_", season_name).strip("_")
    date_str = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"Nhat_Ky_Canh_Tac_{safe}_{date_str}.docx"
    return buf, filename


def export_cultivation_diary_to_docx_file(
    output_path: str,
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> str:
    """Xuất nhật ký canh tác ra file Word (.docx) trên đĩa cứng."""
    data = get_cultivation_diary_export_data(season_id=season_id, farmer_id=farmer_id)
    doc = build_cultivation_diary_docx(data)
    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    doc.save(output_path)
    return output_path


# ==============================================================================
# 3. XUẤT PDF (.PDF)
# ==============================================================================

def build_cultivation_diary_pdf(data: Dict[str, Any], output_target: Union[str, io.BytesIO]):
    """Tạo file PDF nhật ký canh tác hoàn chỉnh."""
    _init_fonts()
    pagesize = pagesizes.landscape(pagesizes.A4)
    margin = 36
    printable_width = pagesize[0] - 2 * margin

    doc = SimpleDocTemplate(
        output_target,
        pagesize=pagesize,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
    )

    styles = getSampleStyleSheet()
    style_title = ParagraphStyle("DTitle", fontName="VN-Times-Bold", fontSize=15, leading=18, alignment=1, textColor=colors.HexColor("#1B5E20"), spaceAfter=10)
    style_meta_val = ParagraphStyle("DMVal", fontName="VN-Times", fontSize=9.5, leading=13, textColor=colors.HexColor("#222222"))
    style_note = ParagraphStyle("DNote", fontName="VN-Times-Italic", fontSize=9, leading=12, alignment=2, textColor=colors.HexColor("#555555"), spaceAfter=6)
    style_th = ParagraphStyle("DTH", fontName="VN-Times-Bold", fontSize=9, leading=11, alignment=1, textColor=colors.HexColor("#1B5E20"))
    style_td = ParagraphStyle("DTD", fontName="VN-Times", fontSize=8.5, leading=11, textColor=colors.HexColor("#212121"))
    style_td_center = ParagraphStyle("DTDCenter", parent=style_td, alignment=1)
    style_td_right = ParagraphStyle("DTDRight", parent=style_td, alignment=2)

    gen = data.get("general_info", {})
    elements = []

    elements.append(Paragraph("SỔ NHẬT KÝ CANH TÁC/Farm diary", style_title))

    meta_pairs = [
        ("Vụ canh tác/Crop season:", gen.get("crop_season", "-")),
        ("Tên hộ sản xuất/Farmer’s name/Supervisor:", gen.get("farmer_name", "-")),
        ("Địa chỉ/Address:", gen.get("address", "-")),
        ("Mã số vùng trồng/P.U.C:", gen.get("farm_code", "-")),
        ("Diện tích canh tác/Planted area:", gen.get("planted_area", "-")),
        ("Giống/Variety:", gen.get("variety", "-")),
        ("Thời điểm bắt đầu trồng:", gen.get("start_date", "-")),
        ("Thời điểm bắt đầu thu hoạch/ Beginning of season:", gen.get("harvest_start_date", "-")),
        ("Sản lượng dự kiến (estimated productivity):", gen.get("expected_yield", "-")),
        ("Sản lượng thực tế (actual productivity):", gen.get("actual_yield", "-")),
        ("Người thu mua (Buyer):", gen.get("buyer", "-")),
    ]

    half = (len(meta_pairs) + 1) // 2
    meta_left = meta_pairs[:half]
    meta_right = meta_pairs[half:]

    meta_table_data = []
    max_len = max(len(meta_left), len(meta_right))
    for i in range(max_len):
        row_cells = []
        if i < len(meta_left):
            lbl, val = meta_left[i]
            row_cells.append(Paragraph(f"<b>{lbl}</b> {val}", style_meta_val))
        else:
            row_cells.append(Paragraph("", style_meta_val))

        if i < len(meta_right):
            lbl, val = meta_right[i]
            row_cells.append(Paragraph(f"<b>{lbl}</b> {val}", style_meta_val))
        else:
            row_cells.append(Paragraph("", style_meta_val))
        meta_table_data.append(row_cells)

    w_half = printable_width / 2.0
    meta_tbl = Table(meta_table_data, colWidths=[w_half, w_half])
    meta_tbl.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    elements.append(meta_tbl)

    elements.append(Paragraph("Ghi chú: nội dung nào không có hoạt động thì ghi là “không”", style_note))

    headers = [
        "Ngày/tháng/năm",
        "Giai đoạn sinh trưởng\ncủa cây trồng",
        "Công việc",
        "Sinh vật gây hại\nphát hiện",
        "Tên phân/thuốc BVTV",
        "Lượng sử dụng",
        "Nơi mua",
        "Thời gian cách ly",
    ]
    col_widths = [65, 95, 165, 90, 100, 75, 105, 74]

    table_data = [[Paragraph(h.replace("\n", "<br/>"), style_th) for h in headers]]
    daily_records = data.get("daily_records", [])

    def _clean_val(v: Any) -> str:
        s = str(v).strip() if v is not None else ""
        return "Không" if s in ("", "-") else s

    for r in daily_records:
        row = [
            Paragraph(r.get("date", ""), style_td_center),
            Paragraph(_clean_val(r.get("growth_stage")), style_td),
            Paragraph(_clean_val(r.get("work_description")), style_td),
            Paragraph(_clean_val(r.get("pest_detected")), style_td),
            Paragraph(_clean_val(r.get("chemical_name")), style_td),
            Paragraph(_clean_val(r.get("quantity_used")), style_td_right),
            Paragraph(_clean_val(r.get("where_bought")), style_td),
            Paragraph(_clean_val(r.get("isolation_period")), style_td_center),
        ]
        table_data.append(row)

    main_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    tbl_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F5E9")),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#AAAAAA")),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for idx in range(1, len(table_data)):
        if idx % 2 == 0:
            tbl_style.append(("BACKGROUND", (0, idx), (-1, idx), colors.HexColor("#F9FBF7")))

    main_table.setStyle(TableStyle(tbl_style))
    elements.append(main_table)

    # Khung ký
    elements.append(Spacer(1, 12))
    now_dt = datetime.now()
    sig_date_str = f"<i>Ngày {now_dt.day:02d} tháng {now_dt.month:02d} năm {now_dt.year}</i>"
    style_sig_title = ParagraphStyle("SigTitle", fontName="VN-Times-Bold", fontSize=9.5, alignment=1, leading=12)
    style_sig_note = ParagraphStyle("SigNote", fontName="VN-Times-Italic", fontSize=8.5, alignment=1, leading=11, textColor=colors.HexColor("#666666"))

    sig_col_w = printable_width / 2.0
    sig_data = [
        ["", Paragraph(sig_date_str, style_sig_note)],
        [Paragraph("NGƯỜI GHI CHÉP / CHỦ HỘ SẢN XUẤT", style_sig_title), Paragraph("KỸ THUẬT VIÊN / CÁN BỘ GIÁM SÁT VIETGAP", style_sig_title)],
        [Paragraph("(Ký và ghi rõ họ tên)", style_sig_note), Paragraph("(Ký, đóng dấu hoặc xác nhận)", style_sig_note)],
        ["", ""],
        ["", ""],
        [Paragraph(f"<b>{gen.get('farmer_name', '')}</b>", style_sig_title), ""],
    ]
    sig_tbl = Table(sig_data, colWidths=[sig_col_w, sig_col_w], rowHeights=[14, 14, 12, 18, 18, 14])
    sig_tbl.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
    elements.append(KeepTogether([sig_tbl]))

    doc.build(elements, canvasmaker=NumberedCanvas)


def export_cultivation_diary_to_pdf_bytes(
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> Tuple[io.BytesIO, str]:
    """Xuất nhật ký canh tác ra buffer PDF."""
    data = get_cultivation_diary_export_data(season_id=season_id, farmer_id=farmer_id)
    buf = io.BytesIO()
    build_cultivation_diary_pdf(data, buf)
    buf.seek(0)
    season_name = data.get("general_info", {}).get("crop_season", "Vu_Canh_Tac")
    safe = re.sub(r"[^\w\-_]+", "_", season_name).strip("_")
    date_str = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"Nhat_Ky_Canh_Tac_{safe}_{date_str}.pdf"
    return buf, filename


def export_cultivation_diary_to_pdf_file(
    output_path: str,
    season_id: Optional[int] = None,
    farmer_id: Optional[Any] = None,
) -> str:
    """Xuất nhật ký canh tác ra file PDF trên đĩa cứng."""
    data = get_cultivation_diary_export_data(season_id=season_id, farmer_id=farmer_id)
    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    build_cultivation_diary_pdf(data, output_path)
    return output_path


# ==============================================================================
# CLI EXECUTION
# ==============================================================================
# Aliases for API flexibility
build_cultivation_diary_excel_workbook = build_cultivation_diary_workbook
export_cultivation_diary_to_excel_bytes = export_cultivation_diary_to_bytes
export_cultivation_diary_to_excel_file = export_cultivation_diary_to_file
export_cultivation_diary_to_excel = export_cultivation_diary_to_file
export_cultivation_diary_to_docx = export_cultivation_diary_to_docx_file
export_cultivation_diary_to_pdf = export_cultivation_diary_to_pdf_file
build_cultivation_diary_docx_document = build_cultivation_diary_docx
build_cultivation_diary_pdf_document = build_cultivation_diary_pdf


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Export Cultivation Diary (Excel, Word, PDF)")
    parser.add_argument("--format", choices=["excel", "docx", "pdf", "all"], default="all", help="Định dạng xuất")
    parser.add_argument("--season-id", type=int, default=None, help="ID vụ mùa")
    parser.add_argument("--farmer-id", type=int, default=None, help="ID nông dân")
    parser.add_argument("--prefix", type=str, default="nhat_ky_canh_tac", help="Tiền tố file")
    args = parser.parse_args()

    fmt = args.format.lower()
    if fmt in ("excel", "all"):
        f_x = f"{args.prefix}.xlsx"
        export_cultivation_diary_to_file(f_x, season_id=args.season_id, farmer_id=args.farmer_id)
        print(f"✅ Đã tạo file Excel: {f_x}")
    if fmt in ("docx", "all"):
        f_d = f"{args.prefix}.docx"
        export_cultivation_diary_to_docx_file(f_d, season_id=args.season_id, farmer_id=args.farmer_id)
        print(f"✅ Đã tạo file Word DOCX: {f_d}")
    if fmt in ("pdf", "all"):
        f_p = f"{args.prefix}.pdf"
        export_cultivation_diary_to_pdf_file(f_p, season_id=args.season_id, farmer_id=args.farmer_id)
        print(f"✅ Đã tạo file PDF: {f_p}")
